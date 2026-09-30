# -*- coding: utf-8 -*-
"""Excel 模板导入器（模板收集模式）。

读「企业材料收集表.xlsx」8 个 sheet + 可选「证明材料/」文件夹，构建 facts
（8 类本体结构）。导入端直接把 facts upsert 进 M3（m3.sync.sync_facts_to_m3），
本导入器只负责解析 ZIP + 落证明材料文件，不再写本地 JSON。

输入：
  - ZIP 包（含一个 .xlsx + 可选 证明材料/ 目录）
  - 或单个 .xlsx 文件（proof_files 为空则无证据路径）

输出：dict = {"facts": {...}, "stats": {...}, "warnings": [...]}
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..document_recognition.storage import DEFAULT_ORGANIZATION_NAME

# ---------------------------------------------------------------------------
# sheet 定义 + 表头 → 字段映射
# ---------------------------------------------------------------------------

# 企业基本信息（竖排「字段 / 填写值」）
ORG_FIELD_MAP: dict[str, str] = {
    "企业全称": "organizationName", "企业类型": "organizationType",
    "统一社会信用代码": "unifiedSocialCreditCode", "法定代表人": "legalRepresentative",
    "注册地址": "registeredAddress", "联系方式": "contactInfo", "官网": "officialSite",
    "从业人员数": "employeeCount", "营业收入(万元)": "annualRevenueWan",
    "营业收入（万元）": "annualRevenueWan", "资产总额(万元)": "totalAssetsWan",
    "资产总额（万元）": "totalAssetsWan", "企业规模": "enterpriseSize",
    "企业描述": "description", "能力清单": "capabilities", "服务范围": "serviceScope",
    "产品范围": "productScope", "区域范围": "regionScope", "行业范围": "industryScope",
    "管理制度/承诺": "managementPolicies", "管理制度／承诺": "managementPolicies",
}

# 事实表 sheet 定义
SHEET_DEFS: list[dict[str, Any]] = [
    {
        "sheet": "资质认证", "key": "qualification_certificate", "id_prefix": "qc",
        "map": {
            "资质类型": "qualificationType", "资质名称": "qualificationName",
            "证书编号": "certificateNo", "等级": "level", "发证机构": "issuer",
            "认证范围": "scope", "颁发日期": "issueDate", "有效期起": "validFrom",
            "有效期止": "validUntil", "状态": "status",
        },
        "numeric": set(),
        "name_field": "qualificationName",
    },
    {
        "sheet": "类似业绩", "key": "performance_case", "id_prefix": "perf",
        "map": {
            "案例名称": "caseName", "客户名称": "customerName", "行业": "industry",
            "地区": "region", "服务范围": "serviceScope", "合同金额(元)": "contractAmount",
            "合同金额（元）": "contractAmount", "币种": "currency", "签订日期": "signedDate",
            "完成日期": "completionDate", "验收状态": "acceptanceStatus", "角色": "role",
            "联系人": "contactName",
        },
        "numeric": {"contractAmount"},
        "name_field": "caseName",
    },
    {
        "sheet": "项目人员", "key": "staff_member", "id_prefix": "staff",
        "map": {
            "姓名": "name", "岗位": "role", "人员类型": "staffType", "职称": "title",
            "专业证书": "professionalCertificates", "从业年限": "yearsOfExperience",
            "入职日期": "employmentStartDate", "专业特长": "specialty",
            "在职状态": "currentStatus", "简历摘要": "resumeSummary",
        },
        "numeric": {"yearsOfExperience"},
        "name_field": "name",
    },
    {
        "sheet": "财务证明", "key": "financial_evidence", "id_prefix": "fin",
        "map": {
            "证明类型": "evidenceType", "期间起": "periodStart", "期间止": "periodEnd",
            "出具机构": "issuedBy", "出具日期": "issueDate", "有效期止": "validUntil",
            "营收(元)": "revenue", "营收（元）": "revenue", "总资产(元)": "assetAmount",
            "总资产（元）": "assetAmount", "信用等级": "creditLevel", "状态": "status",
        },
        "numeric": {"revenue", "assetAmount"},
        "name_field": "evidenceType",
    },
    {
        "sheet": "技术产品", "key": "technical_product", "id_prefix": "prd",
        "map": {
            "产品名称": "productName", "别名": "aliases", "产品类型": "productType",
            "分类": "category", "品牌": "brand", "型号": "model", "描述": "description",
            "适用场景": "applicableScenario",
        },
        "numeric": set(),
        "name_field": "productName",
    },
    {
        "sheet": "技术参数", "key": "technical_parameter", "id_prefix": "param",
        "map": {"参数描述": "description", "所属产品名": "productId"},
        "numeric": set(),
        "name_field": "description",
    },
    {
        "sheet": "证照凭证", "key": "credential", "id_prefix": "cred",
        "map": {
            "凭证类型": "credentialType", "凭证名称": "credentialName",
            "登记号": "registrationNo", "发证机构": "issuer", "权利人": "owner",
            "标准号": "standardNo", "取得日期": "issueDate", "有效期止": "validUntil",
            "关联产品名": "productId",
        },
        "numeric": set(),
        "name_field": "credentialName",
    },
]

_SHEET_BY_KEY: dict[str, dict[str, Any]] = {d["key"]: d for d in SHEET_DEFS}
_SHEET_BY_NAME: dict[str, dict[str, Any]] = {d["sheet"]: d for d in SHEET_DEFS}

# 需要把「产品名」解析成 productId 的 sheet
_PRODUCT_NAME_SHEETS = {"技术参数": "productId", "证照凭证": "productId"}


def _now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _strip_header(h: str) -> str:
    return (h or "").replace("▼", "").strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _to_num(value: Any, kind: str) -> Any:
    if value is None or value == "":
        return None
    s = str(value).strip().replace(",", "").replace("，", "")
    # 去掉常见单位后缀（万元/元/万/%/人 等），避免 "24049.15万元" 写不进 M3 的 double 字段
    s = re.sub(r"(万元|元|万|%|人|个|家)$", "", s).strip()
    if not s:
        return None
    try:
        if kind == "int":
            return int(float(s))
        return float(s)
    except ValueError:
        return value


def _norm_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        # Excel 日期序列号（1900 纪元）→ yyyy-MM-dd（如 43469 → 2019-01-04）
        try:
            if 1 <= value <= 2958465:
                return (datetime(1899, 12, 30) + timedelta(days=int(value))).strftime("%Y-%m-%d")
        except Exception:
            pass
        return str(int(value))
    s = str(value).strip()
    # 纯数字字符串（可能是 Excel 日期序列号，如 "43469"）→ 转成 yyyy-MM-dd
    if re.fullmatch(r"\d{3,7}", s):
        try:
            n = int(s)
            if 1 <= n <= 2958465:
                return (datetime(1899, 12, 30) + timedelta(days=n)).strftime("%Y-%m-%d")
        except Exception:
            pass
        return ""
    # 永久有效语义 → 哨兵日期 9999-12-31（M3 date 字段可比较，匹配引擎直接判"未过期"）
    if s in ("长期", "长期有效", "任期内有效", "每三年一复核"):
        return "9999-12-31"
    # 无有效期约束 → 置空
    if s in ("-", "--", "无"):
        return ""
    # 相对有效期（"三年"）→ 标记，由 _parse_table 按颁发日期 +3 年换算
    if s == "三年":
        return "__PLUS_3Y__"
    m = re.search(r"(\d{4})[年/\-.](\d{1,2})[月/\-.](\d{1,2})日?", s)
    if m:
        year, month, day = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"{year:04d}-{month:02d}-{day:02d}"
    # 兜底：无法解析成标准日期的文本 → 置空
    return ""


def _date_plus_years(base: str, years: int) -> str:
    """基准日期 + N 年，返回 yyyy-MM-dd；基准缺失/非法时按永久有效兜底。"""
    if not base or base in ("", "__PLUS_3Y__"):
        return "9999-12-31"
    try:
        d = datetime.strptime(base.split("T")[0].split(" ")[0], "%Y-%m-%d")
        try:
            return d.replace(year=d.year + years).strftime("%Y-%m-%d")
        except ValueError:
            # 2/29 平年无此日 → 退到 2/28
            return d.replace(year=d.year + years, month=2, day=28).strftime("%Y-%m-%d")
    except ValueError:
        return "9999-12-31"


def _norm_seq(s: str) -> str:
    """归一化序号：纯数字去前导零（'003'→'3'），其余保留 strip 原文。"""
    s = (s or "").strip()
    return str(int(s)) if re.fullmatch(r"\d+", s) else s


def _norm_name(s: str) -> str:
    """归一化产品名：去空白 + 小写，用于容错匹配（忽略空格/大小写差异）。"""
    return re.sub(r"\s+", "", str(s or "")).lower()


def _split_aliases(value: str) -> list[str]:
    """拆产品别名（中英文逗号/顿号/分号/斜杠/空白）。"""
    return [a.strip() for a in re.split(r"[,，、;；/\s]+", str(value or "")) if a.strip()]


def _parse_proof_arcname(arcname: str) -> tuple[str, str, str] | None:
    """解析证明材料路径 → (sheet, 序号, 文件名)。

    新规范：证明材料/{sheet}/{序号}/{文件名}（一个序号一个文件夹，N 个文件）
    旧平铺：证明材料/{sheet}/{序号}.ext（兼容，序号取自文件名 stem）
    """
    parts = arcname.replace("\\", "/").strip("/").split("/")
    if len(parts) < 3 or parts[0] != "证明材料":
        return None
    sheet = parts[1]
    if len(parts) >= 4:
        seq = parts[2]
        filename = parts[-1]
    else:
        filename = parts[-1]
        seq = Path(filename).stem
    return sheet, seq, filename


class ExcelImporter:
    def __init__(self, runtime_dir: Path | None = None):
        self.runtime_dir = Path(runtime_dir) if runtime_dir else None

    # ---------------------------------------------------------- 入口
    def import_zip(self, zip_bytes: bytes) -> dict[str, Any]:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            members = zf.namelist()
            xlsx_name = next(
                (n for n in members if n.lower().endswith(".xlsx") and "__MACOSX" not in n), None
            )
            if not xlsx_name:
                raise ValueError("ZIP 里没有 .xlsx 文件")
            xlsx_bytes = zf.read(xlsx_name)
            proof_files: dict[str, bytes] = {}
            for n in members:
                if n.startswith("证明材料/") and not n.endswith("/") and "__MACOSX" not in n:
                    proof_files[n] = zf.read(n)
            return self.import_xlsx(xlsx_bytes, proof_files=proof_files)

    def import_xlsx(
        self, xlsx_bytes: bytes, proof_files: dict[str, bytes] | None = None
    ) -> dict[str, Any]:
        proof_files = proof_files or {}
        wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        available = set(wb.sheetnames)

        warnings: list[str] = []
        facts: dict[str, Any] = {"organization": None}
        stats: dict[str, int] = {}

        # 1) 企业基本信息
        if "企业基本信息" in available:
            facts["organization"] = self._parse_organization(wb["企业基本信息"])
        else:
            warnings.append("缺少「企业基本信息」sheet，企业主体用默认名")

        # 2) 各事实表
        for d in SHEET_DEFS:
            key = d["key"]
            facts[key] = []
            if d["sheet"] in available:
                rows = self._parse_table(wb[d["sheet"]], d)
                for row in rows:
                    if any(v not in (None, "") for k, v in row.items() if k != "__seq__"):
                        facts[key].append(row)
            else:
                warnings.append(f"缺少「{d['sheet']}」sheet，按空表处理（现有工作表：{', '.join(available)}）")
            stats[key] = len(facts[key])

        # 3) 生成主键 ID（必须先于产品名解析：技术产品要有 productId，参数/证照才能引用）
        self._assign_ids(facts)

        # 4) 解析产品名 → productId（把技术参数/证照凭证里的「产品名」换成产品主键）
        #    匹配策略：精确 → 归一化（去空格/忽略大小写，含别名）；命中非精确时记 warning 便于溯源。
        product_name_to_id = {
            p.get("productName"): p.get("productId")
            for p in facts.get("technical_product", []) if p.get("productName")
        }
        norm_to_ids: dict[str, list[tuple[str, str]]] = {}
        for p in facts.get("technical_product", []):
            pname = p.get("productName")
            pid = p.get("productId")
            if not pname or not pid:
                continue
            tokens = [pname] + _split_aliases(p.get("aliases"))
            for token in tokens:
                norm_to_ids.setdefault(_norm_name(token), []).append((pname, pid))
        for sheet_name, field in _PRODUCT_NAME_SHEETS.items():
            key = _SHEET_BY_NAME[sheet_name]["key"]
            for row in facts.get(key, []):
                name = row.get(field)
                if not name:
                    continue
                if name in product_name_to_id:
                    row[field] = product_name_to_id[name]
                    continue
                candidates = norm_to_ids.get(_norm_name(name), [])
                if not candidates:
                    warnings.append(f"「{sheet_name}」里填的产品名「{name}」在「技术产品」表未找到，未关联")
                    row[field] = ""
                elif len(candidates) > 1:
                    warnings.append(
                        f"「{sheet_name}」里填的产品名「{name}」归一化后匹配到 {len(candidates)} 个产品"
                        f"（{', '.join(c[0] for c in candidates[:3])}…），无法唯一确定，未关联"
                    )
                    row[field] = ""
                else:
                    row[field] = candidates[0][1]
                    warnings.append(
                        f"「{sheet_name}」里填的产品名「{name}」→ 归一化匹配到「{candidates[0][0]}」"
                    )

        # 5) 提取证明材料文件到 runtime/libraries/company/
        self._extract_proof_files(proof_files)

        # 6) 关联证据文件夹路径（一个序号一个文件夹，只记路径不逐文件建材料）
        evidence_count = self._assign_evidence(facts, proof_files, warnings)
        stats["evidence_paths"] = evidence_count

        # 7) 清理内部字段
        self._strip_internal(facts)

        return {"facts": facts, "stats": stats, "warnings": warnings}

    # ---------------------------------------------------------- 解析
    def _parse_organization(self, ws) -> dict[str, Any]:
        values: dict[str, str] = {}
        for row in ws.iter_rows(min_row=1, max_col=3):
            if len(row) < 2:
                continue
            label = _cell_text(row[0].value)
            val = _cell_text(row[1].value)
            if label:
                values[label] = val

        def _get(*keys: str) -> str:
            for k in keys:
                if values.get(k):
                    return values[k]
            return ""

        org = {
            "organizationId": "org-001",
            "organizationName": _get("企业全称") or DEFAULT_ORGANIZATION_NAME,
            "organizationType": _get("企业类型") or "bidder",
            "unifiedSocialCreditCode": _get("统一社会信用代码"),
            "legalRepresentative": _get("法定代表人"),
            "registeredAddress": _get("注册地址"),
            "contactInfo": _get("联系方式"),
            "officialSite": _get("官网"),
            "aliases": "",
            "employeeCount": _to_num(_get("从业人员数"), "int"),
            "annualRevenueWan": _to_num(_get("营业收入(万元)", "营业收入（万元）"), "float"),
            "totalAssetsWan": _to_num(_get("资产总额(万元)", "资产总额（万元）"), "float"),
            "enterpriseSize": _get("企业规模"),
            "description": _get("企业描述"),
            "managementPolicies": _get("管理制度/承诺", "管理制度／承诺"),
            "capabilities": _get("能力清单"),
            "serviceScope": _get("服务范围"),
            "productScope": _get("产品范围"),
            "regionScope": _get("区域范围"),
            "industryScope": _get("行业范围"),
            "scaleValue": 0,
            "scaleUnit": "",
            "status": "active",
        }
        for f in ("employeeCount", "annualRevenueWan", "totalAssetsWan"):
            if org[f] in (0, ""):
                org[f] = None
        return org

    def _parse_table(self, ws, d: dict[str, Any]) -> list[dict[str, Any]]:
        col_map: dict[str, str] = d["map"]
        numeric: set[str] = d["numeric"]
        date_fields = {"issueDate", "validFrom", "validUntil", "signedDate",
                       "completionDate", "periodStart", "periodEnd", "employmentStartDate"}

        headers: list[str] = []
        for row in ws.iter_rows(min_row=1, max_row=1):
            headers = [_strip_header(_cell_text(c.value)) for c in row]
            break

        col_to_field: dict[int, str] = {}
        for idx, h in enumerate(headers):
            if h == "序号":
                col_to_field[idx] = "__seq__"
            elif h in col_map:
                col_to_field[idx] = col_map[h]

        rows: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=2):
            item: dict[str, Any] = {}
            for idx, field in col_to_field.items():
                if idx >= len(row):
                    continue
                raw = row[idx].value
                if field == "__seq__":
                    item[field] = _cell_text(raw)
                elif field in numeric:
                    item[field] = _to_num(raw, "float")
                elif field in date_fields:
                    item[field] = _norm_date(raw)
                else:
                    item[field] = _cell_text(raw)
            # 相对有效期（"三年"）→ 按颁发/签订日期 +3 年换算成具体日期
            for f in date_fields:
                if item.get(f) == "__PLUS_3Y__":
                    base = (item.get("issueDate") or item.get("signedDate")
                            or item.get("validFrom") or item.get("periodStart") or "")
                    item[f] = _date_plus_years(base, 3)
            if any(v not in (None, "") for k, v in item.items() if k != "__seq__"):
                rows.append(item)
        return rows

    # ---------------------------------------------------------- 主键
    def _assign_ids(self, facts: dict[str, Any]) -> None:
        for key, d in _SHEET_BY_KEY.items():
            id_field = {
                "qualification_certificate": "qualificationId",
                "performance_case": "performanceId",
                "staff_member": "staffId",
                "financial_evidence": "financialEvidenceId",
                "technical_product": "productId",
                "technical_parameter": "parameterId",
                "credential": "credentialId",
            }[key]
            for i, row in enumerate(facts.get(key, [])):
                # 序号优先用 Excel「序号」列（归一化去前导零），无序号列时回退到行号。
                # 主键 = 「前缀-序号」（不带随机批次），这样同一主表反复导入时 M3 按
                # biz_id upsert 而不产生重复对象；序号必须按「追加式」维护（用了不重排、
                # 删了留空洞、新增 = max+1），保证主键 ID、Excel 序号列、证明材料文件夹名三者永远同源。
                seq = _norm_seq(row.get("__seq__") or "") or str(i + 1)
                row[id_field] = f"{d['id_prefix']}-{seq}"

    # ---------------------------------------------------------- 证据文件夹
    def _extract_proof_files(self, proof_files: dict[str, bytes]) -> None:
        """把证明材料文件写入 runtime/libraries/company/{arcname}（相对路径可被 /runtime 服务）。"""
        if not self.runtime_dir:
            return
        base = (self.runtime_dir / "libraries" / "company").resolve()
        for arcname, content in proof_files.items():
            rel = arcname.replace("\\", "/").strip("/")
            if not rel or rel.startswith("../") or "/../" in rel:
                continue
            target = (base / rel).resolve()
            if base != target and base not in target.parents:
                continue  # 路径穿越，跳过
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(content)

    def _assign_evidence(
        self, facts: dict[str, Any], proof_files: dict[str, bytes], warnings: list[str]
    ) -> int:
        """把证明材料文件夹路径写回事实的 evidencePath（不逐文件建材料）。

        一个序号一个文件夹：evidencePath = libraries/company/证明材料/{sheet}/{序号}/
        （相对 runtime_root，文件由 /runtime 静态路由 + /api/materials/evidence-files 提供）。
        返回已关联证据的事实数。
        """
        # (sheet, 归一化序号) → fact
        fact_index: dict[tuple[str, str], dict[str, Any]] = {}
        for d in SHEET_DEFS:
            for row in facts.get(d["key"], []):
                seq = _norm_seq(row.get("__seq__", ""))
                if seq:
                    fact_index[(d["sheet"], seq)] = row

        # 默认：所有事实证据路径为空，并挂企业主键
        for d in SHEET_DEFS:
            for row in facts.get(d["key"], []):
                row["organizationId"] = "org-001"
                row.setdefault("evidencePath", "")
                row.setdefault("evidenceUpdatedAt", "")

        # (sheet, 序号) → 文件夹相对路径（取每个文件所在目录）
        folders: dict[tuple[str, str], str] = {}
        for arcname in proof_files:
            parsed = _parse_proof_arcname(arcname)
            if not parsed:
                continue
            sheet, seq_raw, _ = parsed
            seq = _norm_seq(seq_raw)
            folder = arcname.replace("\\", "/").strip("/").rsplit("/", 1)[0] + "/"
            folders.setdefault((sheet, seq), "libraries/company/" + folder)

        count = 0
        for (sheet, seq), folder in sorted(folders.items()):
            fact = fact_index.get((sheet, seq))
            if fact is None:
                warnings.append(f"证明材料/{sheet}/{seq}/ 没有对应的表行，未关联")
                continue
            fact["evidencePath"] = folder
            fact["evidenceUpdatedAt"] = _now_iso()
            count += 1
        return count

    @staticmethod
    def _strip_internal(facts: dict[str, Any]) -> None:
        for key in _SHEET_BY_KEY:
            for row in facts.get(key, []):
                row.pop("__seq__", None)
