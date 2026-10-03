# -*- coding: utf-8 -*-
"""Deterministic round-trip tests for template, response-table, and evidence flows."""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.app.bid_writing.assemble import assemble_docx, render_template_preview_docx
from backend.app.bid_writing.evidence_chapter import build_evidence_chapter_result
from backend.app.bid_writing.response_table import (
    build_response_table_result,
    build_template_response_table_spec,
    prepare_response_table_inputs,
)
from backend.app.bid_writing.template_reuse import fill_template_fragment
from backend.tests.fixtures.fidelity_fixtures import (
    evidence_fixture,
    response_table_inputs,
    template_bundle,
)


ROOT = Path(__file__).resolve().parents[2]
BASELINE_GENERATOR = ROOT / "backend" / "tests" / "baselines" / "generate_fidelity_baseline.py"


class FidelityRoundTripTests(unittest.TestCase):
    def test_template_roundtrip_preserves_structure_style_and_field_allowlist(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fixture = template_bundle(root)
            fragment = fixture["fragment"]

            self.assertEqual(fragment["bindingStatus"], "ready")
            self.assertEqual(fragment["copyMode"], "ooxml_body_range")
            for fixed_text in ("投标函格式", "后续固定内容", "锁定文字"):
                self.assertIn(fixed_text, fragment["templateText"])

            table = fragment["tables"][0]
            self.assertEqual((table["rowCount"], table["columnCount"]), (3, 3))
            cells = [
                cell
                for row in table["rows"]
                for cell in row["cells"]
            ]
            self.assertEqual(
                [cell["gridSpan"] for cell in cells],
                ["2", "2", "1", "1", "1", "1", "1", "1", "1"],
            )
            self.assertTrue(all(cell["vMerge"] == "" for cell in cells))

            heading = next(block for block in fragment["blocks"] if block["text"] == "投标函格式")
            self.assertEqual(heading["style"]["styleId"], "Heading2")
            self.assertEqual(heading["style"]["bold"], True)
            self.assertEqual(heading["style"]["fontSize"], 16.0)
            self.assertIn("CENTER", heading["style"]["alignment"])

            allowed_fields = {"field-project", "field-company"}
            self.assertEqual(
                {field["fieldId"] for field in fragment["editableFields"]},
                allowed_fields,
            )

            filled = fill_template_fragment(
                {"sectionId": "section-fixture", "title": "投标函"},
                fragment,
                {"parsed": {"projectName": "测试项目"}},
                {"organizationName": "测试企业"},
                {},
                agent_result={
                    "fieldValues": [
                        {
                            "fieldId": "field-company",
                            "label": "投标人名称",
                            "value": "测试企业",
                            "status": "filled",
                            "sourceRefs": [{"factKey": "organization", "name": "企业名称"}],
                        },
                        {
                            "fieldId": "unknown-field",
                            "label": "未授权字段",
                            "value": "不应填入",
                            "status": "filled",
                            "sourceRefs": [{"factKey": "unknown", "name": "未知"}],
                        },
                    ]
                },
            )
            self.assertIn("项目名称：测试项目", filled["content"])
            self.assertIn("投标人名称：测试企业", filled["content"])
            self.assertNotIn("不应填入", filled["content"])
            self.assertTrue(
                {field["fieldId"] for field in filled["templateFieldValues"]}
                <= allowed_fields
            )

            section = {
                "sectionId": "section-fixture",
                "title": "投标函",
                "source": "copied_template",
                "templateFragmentId": "tpl-fixture",
                "content": filled["content"],
                "contentMatchesTemplate": True,
                "templateFieldValues": filled["templateFieldValues"],
                "templateBlocks": filled["templateBlocks"],
                "templateTables": filled["templateTables"],
            }
            output = assemble_docx(
                {
                    "title": "测试项目",
                    "mode": "open",
                    "plan": [section],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "output",
                root=root,
                template_fragments=[fragment],
            )
            self.assertEqual(output["copiedTemplateCount"], 1)
            exported = Document(output["path"])
            exported_heading = next(
                paragraph for paragraph in exported.paragraphs if paragraph.text == "投标函格式"
            )
            self.assertEqual(exported_heading.alignment, 1)
            self.assertTrue(exported_heading.runs[0].bold)
            self.assertEqual(exported_heading.runs[0].font.size.pt, 16.0)
            self.assertIn(
                "锁定文字",
                "\n".join(cell.text for table in exported.tables for row in table.rows for cell in row.cells),
            )

            preview = render_template_preview_docx(fragment, section, root=root)
            self.assertIsNotNone(preview)
            assert preview is not None
            preview_document = Document(__import__("io").BytesIO(preview[0]))
            preview_heading = next(
                paragraph for paragraph in preview_document.paragraphs if paragraph.text == "投标函格式"
            )
            self.assertEqual(preview_heading.alignment, 1)
            self.assertTrue(preview_heading.runs[0].bold)
            self.assertEqual(preview_heading.runs[0].font.size.pt, 16.0)
            self.assertEqual(len(preview_document.tables[0].rows), 3)
            self.assertEqual(len(preview_document.tables[0].columns), 3)

    def test_response_table_roundtrip_preserves_rows_status_and_sources(self) -> None:
        section, fine_review, knowledge = response_table_inputs()
        prepared = prepare_response_table_inputs(section, fine_review, knowledge)
        self.assertEqual(len(prepared), 2)
        self.assertEqual([row["requirementId"] for row in prepared], ["REQ-001", "REQ-002"])

        result = build_response_table_result(
            section,
            prepared,
            {"REQ-001": "支持 SSO 单点登录。"},
        )
        rows = result["requirementResponseRows"]
        self.assertEqual(len(rows), 2)
        self.assertEqual([row["requirementId"] for row in rows], ["REQ-001", "REQ-002"])
        self.assertEqual([row["matchStatus"] for row in rows], ["满足", "部分满足"])
        self.assertFalse(rows[0]["needsReview"])
        self.assertTrue(rows[1]["needsReview"])
        self.assertEqual(rows[0]["response"], "支持 SSO 单点登录。")
        self.assertIn("需人工补充", rows[1]["response"])
        self.assertEqual(
            [ref["evidencePath"] for ref in rows[0]["sourceRefs"]],
            ["evidence/product-001.pdf"],
        )
        self.assertEqual(len(result["tableMarkdown"].splitlines()[0].split("|")[1:-1]), 5)

        fragment = {
            "tables": [
                {
                    "tableId": "TABLE-RESPONSE",
                    "rows": [
                        {
                            "rowIndex": 1,
                            "cells": [
                                {"cellIndex": 1, "text": "序号"},
                                {"cellIndex": 2, "text": "技术和服务要求"},
                                {"cellIndex": 3, "text": "投标人响应"},
                                {"cellIndex": 4, "text": "偏离情况"},
                                {"cellIndex": 5, "text": "备注"},
                            ],
                        },
                        {
                            "rowIndex": 2,
                            "cells": [{"cellIndex": index, "text": ""} for index in range(1, 6)],
                        },
                    ],
                }
            ]
        }
        mapped = build_template_response_table_spec(fragment, rows)
        self.assertIsNotNone(mapped)
        assert mapped is not None
        self.assertEqual(mapped["columns"], {"sequence": 1, "requirement": 2, "response": 3, "deviation": 4, "remark": 5})
        self.assertEqual(len(mapped["rows"]), 1)
        self.assertEqual(mapped["rows"][0]["requirementId"], "REQ-001")
        self.assertEqual(len(mapped["appendRows"]), 1)
        self.assertEqual(mapped["appendRows"][0]["requirementId"], "REQ-002")

    def test_evidence_roundtrip_checks_files_order_reuse_and_unreadable_todo(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fixture = evidence_fixture(root)
            refs = fixture["refs"]
            self.assertTrue(Path(refs["readable"]).is_file())
            self.assertTrue(Path(refs["second"]).is_file())
            self.assertFalse(Path(refs["unreadable"]).is_file())

            chapter = build_evidence_chapter_result(
                {"title": "技术参数证明材料"}, fixture["entries"], fixture["assets"]
            )
            self.assertEqual(chapter["summary"]["total"], 3)
            self.assertEqual(chapter["summary"]["ready"], 2)
            self.assertEqual(chapter["summary"]["missing"], 1)
            self.assertEqual(
                [Path(ref["evidencePath"]).name for group in chapter["evidenceGroups"] for ref in group["insertRefs"]],
                ["01-product.pdf", "02-certificate.pdf", "03-missing.pdf"],
            )
            self.assertEqual(
                [Path(ref["evidencePath"]).name for group in chapter["evidenceGroups"] for ref in group["reusedRefs"]],
                ["01-product.pdf"],
            )
            self.assertIn("REQ-003", "\n".join(chapter["blanksToFill"]))
            self.assertFalse(fixture["assets"][refs["unreadable"]]["available"])
            self.assertTrue(fixture["assets"][refs["unreadable"]]["placeholder"])
            self.assertIn("待补扫描件", chapter["content"])

            output = assemble_docx(
                {
                    "title": "证据测试项目",
                    "mode": "open",
                    "plan": [
                        {
                            "sectionId": "evidence-01",
                            "title": "技术参数证明材料",
                            "responseType": "attachment",
                            "content": chapter["content"],
                            "evidenceGroups": chapter["evidenceGroups"],
                            "sourceRefs": chapter["sourceRefs"],
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "output",
                assets=fixture["assets"],
            )
            paragraphs = "\n".join(paragraph.text for paragraph in Document(output["path"]).paragraphs)
            self.assertLess(paragraphs.index("产品证明"), paragraphs.index("证书证明"))
            self.assertLess(paragraphs.index("证书证明"), paragraphs.index("待补证明"))
            self.assertIn("待补扫描件", paragraphs)

    def test_fidelity_baseline_is_reproducible_without_runtime_data(self) -> None:
        result = subprocess.run(
            [sys.executable, str(BASELINE_GENERATOR), "--check"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            cwd=str(ROOT),
        )
        output = (result.stdout or "") + (result.stderr or "")
        self.assertEqual(result.returncode, 0, f"保真基线不一致：\n{output}")


if __name__ == "__main__":
    unittest.main()
