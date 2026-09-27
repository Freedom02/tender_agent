import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.app.bid_writing.assemble import assemble_docx, render_template_preview_docx
from backend.app.bid_writing.template_fragments import (
    attach_template_fragments_to_bundle,
    attach_template_fragments_to_plan,
    match_template_fragment,
)
from backend.app.bid_writing.template_reuse import (
    fill_template_fragment,
    source_table_for_section,
)
from backend.app.pi_agent.writing_orchestrator import PiWritingOrchestrator
from backend.app.store import Store
from backend.app.tender_matching.document_parser import parse_docx
from backend.app.tender_matching.document_structure import enrich_document_structure


class FakeTemplateFillClient:
    def __init__(self) -> None:
        self.config = SimpleNamespace(model="pi-template-test-model")
        self.calls: list[str] = []

    def fill_template_fields(self, **_: object) -> dict[str, object]:
        self.calls.append("TableFormFillAgent")
        return {
            "templateId": "tpl-001",
            "sectionId": "c-01",
            "status": "completed",
            "fieldValues": [
                {
                    "fieldId": "field-001",
                    "label": "投标人名称",
                    "value": "测试企业",
                    "status": "filled",
                    "sourceRefs": [{"factKey": "organization", "name": "企业名称"}],
                }
            ],
            "cellValues": [],
            "blanksToFill": ["投标报价"],
            "warnings": [],
        }


class TemplateFragmentTests(unittest.TestCase):
    def _source_docx(self, root: Path) -> Path:
        source = root / "runtime" / "projects" / "p1" / "source" / "tender.docx"
        source.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading("第一章 投标文件格式", level=1)
        document.add_paragraph("投标函格式")
        document.add_paragraph("投标人名称：________")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "固定表头"
        table.cell(0, 1).text = "填写项"
        table.cell(1, 0).text = "项目名称"
        table.cell(1, 1).text = "________"
        document.add_paragraph("后续固定内容")
        document.save(source)
        return source

    def _bundle(self, root: Path, source: Path) -> dict[str, object]:
        parsed = enrich_document_structure(parse_docx(source.read_bytes()))
        return attach_template_fragments_to_bundle(
            {
                "projectId": "p1",
                "workspaceId": "w1",
                "sourceFile": {
                    "filename": "tender.docx",
                    "storagePath": "runtime/projects/p1/source/tender.docx",
                },
                "documentBlocks": parsed,
                "bidFormatSpec": {
                    "templates": [
                        {
                            "templateId": "tpl-001",
                            "title": "投标函格式",
                            "templateType": "fixed_template",
                            "sourceLocator": "投标函格式",
                        }
                    ]
                },
            },
            root=root,
            source_path=source,
        )

    def test_source_table_matches_technical_response_headers_without_template_fragment(self) -> None:
        unrelated_table = {
            "tableId": "TABLE-PRICE",
            "rows": [
                {
                    "cells": [
                        {"text": "序号"},
                        {"text": "商品名称"},
                        {"text": "投标报价"},
                    ]
                }
            ],
        }
        response_table = {
            "tableId": "TABLE-RESPONSE",
            "rows": [
                {
                    "cells": [
                        {"text": "序号"},
                        {"text": "技术和服务要求"},
                        {"text": "投标人响应"},
                        {"text": "偏离说明"},
                    ]
                }
            ],
        }

        selected = source_table_for_section(
            {"title": "技术和服务要求响应表", "responseType": "table"},
            [unrelated_table, response_table],
        )

        self.assertIs(selected, response_table)

    def test_binds_docx_body_range_and_writes_fragment_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            bundle = self._bundle(root, source)

            fragments = bundle["templateFragments"]
            assert isinstance(fragments, list)
            fragment = fragments[0]
            self.assertEqual(fragment["copyMode"], "ooxml_body_range")
            self.assertGreaterEqual(fragment["bodyChildStart"], 0)
            self.assertGreaterEqual(fragment["bodyChildEnd"], fragment["bodyChildStart"])
            self.assertTrue(fragment["editableFields"])
            self.assertTrue((root / fragment["fragmentArtifact"]).is_file())

            plan = attach_template_fragments_to_plan(
                [{"sectionId": "c-01", "title": "投标函", "source": "template"}],
                fragments,
            )
            self.assertEqual(plan[0]["source"], "copied_template")
            self.assertEqual(plan[0]["templateFragmentId"], "tpl-001")

    def test_submission_structure_creates_fragments_without_pi_template_list(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            parsed = enrich_document_structure(parse_docx(source.read_bytes()))
            bundle = attach_template_fragments_to_bundle(
                {
                    "projectId": "p1",
                    "sourceFile": {"filename": "tender.docx"},
                    "documentBlocks": parsed,
                    "bidFormatSpec": {},
                    "submissionStructure": {
                        "volumes": [
                            {
                                "volumeId": "qualification",
                                "title": "资格及资信证明部分",
                                "items": [
                                    {
                                        "formatItemId": "qualification-01",
                                        "title": "投标函",
                                        "itemType": "fixed_form",
                                        "responseType": "fixed_form",
                                        "templateType": "fixed_template",
                                    }
                                ],
                            }
                        ]
                    },
                },
                root=root,
                source_path=source,
            )
            fragments = bundle["templateFragments"]
            self.assertEqual(len(fragments), 1)
            self.assertEqual(fragments[0]["formatItemId"], "qualification-01")
            self.assertEqual(fragments[0]["copyMode"], "ooxml_body_range")
            self.assertEqual(
                fragments[0]["sourceFile"]["storagePath"],
                source.relative_to(root).as_posix(),
            )

    def test_format_chapter_heading_wins_over_composition_list_and_index(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "runtime" / "projects" / "p1" / "source" / "tender.docx"
            source.parent.mkdir(parents=True, exist_ok=True)
            document = Document()
            document.add_heading("第三章 投标人须知", level=1)
            document.add_paragraph("10.2电子投标文件由下述部分组成：")
            document.add_paragraph("①投标函")
            document.add_heading("第七章 电子投标文件格式", level=1)
            document.add_paragraph("索引")
            document.add_paragraph("一、投标函")
            document.add_paragraph("二、投标人的资格及资信证明文件")
            document.add_heading("一、投标函", level=2)
            document.add_paragraph("投标人名称：________")
            document.add_paragraph("项目名称：________")
            document.add_heading("二、投标人的资格及资信证明文件", level=2)
            document.add_paragraph("资格证明材料")
            document.save(source)

            parsed = enrich_document_structure(parse_docx(source.read_bytes()))
            bundle = attach_template_fragments_to_bundle(
                {
                    "projectId": "p1",
                    "sourceFile": {"filename": "tender.docx"},
                    "documentBlocks": parsed,
                    "submissionStructure": {
                        "volumes": [
                            {
                                "volumeId": "qualification",
                                "title": "资格及资信证明部分",
                                "items": [
                                    {
                                        "formatItemId": "qualification-01",
                                        "title": "投标函",
                                        "itemType": "fixed_form",
                                        "responseType": "fixed_form",
                                        "templateType": "fixed_template",
                                    }
                                ],
                            }
                        ]
                    },
                },
                root=root,
                source_path=source,
            )

            fragment = bundle["templateFragments"][0]
            self.assertEqual(fragment["bindingStatus"], "ready")
            self.assertEqual(fragment["copyMode"], "ooxml_body_range")
            self.assertIn("投标人名称", fragment["templateText"])
            self.assertIn("项目名称", fragment["templateText"])
            self.assertNotIn("资格证明材料", fragment["templateText"])
            self.assertGreater(fragment["bodyChildStart"], 0)

    def test_title_substring_does_not_bind_a_different_section(self) -> None:
        fragment = {
            "templateId": "tpl-001",
            "title": "投标函",
            "bindingStatus": "ready",
            "templateText": "投标函：____",
        }
        self.assertEqual(
            match_template_fragment(
                {"sectionId": "s-1", "title": "投标函说明"},
                [fragment],
            ),
            {},
        )

    def test_sensitive_template_fields_stay_for_confirmation(self) -> None:
        result = fill_template_fragment(
            {"sectionId": "c-01", "title": "投标函"},
            {
                "templateId": "tpl-001",
                "templateText": (
                    "项目名称：________\n"
                    "服务期限：________\n"
                    "服务地点：________\n"
                    "投标报价：________\n"
                    "日期：____年__月__日"
                ),
                "editableFields": [
                    {"fieldId": "f-project", "label": "项目名称", "sourceKey": "projectName"},
                    {"fieldId": "f-period", "label": "服务期限", "sourceKey": "servicePeriod"},
                    {"fieldId": "f-location", "label": "服务地点", "sourceKey": "serviceLocation"},
                    {"fieldId": "f-price", "label": "投标报价", "sourceKey": "bidPrice"},
                ],
            },
            {
                "parsed": {
                    "projectName": "测试项目",
                    "servicePeriod": "12个月",
                    "serviceLocation": "福州",
                    "budgetAmount": "100万元",
                }
            },
            {"organizationName": "本地企业"},
            {},
        )
        self.assertIn("项目名称：测试项目", result["content"])
        self.assertIn("服务期限：________", result["content"])
        self.assertIn("服务地点：________", result["content"])
        self.assertIn("投标报价：________", result["content"])
        self.assertIn("落款日期", result["blanksToFill"])
        self.assertIn("服务期限", result["blanksToFill"])
        self.assertEqual(result["templateFillStatus"], "needs_confirmation")

    def test_fill_template_replaces_government_procurement_fill_markers(self) -> None:
        result = fill_template_fragment(
            {"sectionId": "c-01", "title": "投标函"},
            {
                "templateId": "tpl-001",
                "templateText": (
                    "项目名称：（由投标人填写）\n"
                    "投标人：（填写“全称”）\n"
                    "兹收到贵单位关于（填写“项目名称”）项目（项目编号： ）的投标邀请"
                ),
            },
            {"title": "数字孪生应用实训平台采购", "parsed": {"projectCode": "FJUS-001"}},
            {"organizationName": "测试企业"},
            {},
        )

        self.assertIn("项目名称：数字孪生应用实训平台采购", result["content"])
        self.assertIn("投标人：测试企业", result["content"])
        self.assertIn("关于（数字孪生应用实训平台采购）项目", result["content"])

    def test_ambiguous_template_title_is_not_bound_to_whole_section(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            parsed = enrich_document_structure(parse_docx(source.read_bytes()))
            parsed["segments"].append({
                **parsed["segments"][0],
                "id": "SEG-DUPLICATE",
                "text": "投标函格式",
                "bodyChildIndex": 5,
            })
            parsed["units"] = parsed["segments"]
            for section in parsed.get("sections", []):
                if isinstance(section, dict):
                    section.setdefault("segmentIds", []).append("SEG-DUPLICATE")
            bundle = attach_template_fragments_to_bundle(
                {
                    "projectId": "p1",
                    "sourceFile": {"filename": "tender.docx"},
                    "documentBlocks": parsed,
                    "bidFormatSpec": {
                        "templates": [{
                            "templateId": "tpl-001",
                            "formatItemId": "commercial-01",
                            "title": "投标函格式",
                        }]
                    },
                },
                source_path=source,
            )
            fragment = bundle["templateFragments"][0]
            self.assertEqual(fragment["copyMode"], "text_fallback")
            self.assertTrue(any("唯一" in warning for warning in fragment["warnings"]))

    def test_fill_keeps_missing_fields_and_adds_agent_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            bundle = self._bundle(root, source)
            fragment = bundle["templateFragments"][0]
            result = fill_template_fragment(
                {"sectionId": "c-01", "title": "投标函"},
                fragment,
                {"parsed": {"projectName": "测试项目"}},
                {"organizationName": "本地企业"},
                {},
                agent_result={
                    "fieldValues": [
                        {
                            "fieldId": "field-agent",
                            "label": "投标人名称",
                            "value": "测试企业",
                            "status": "filled",
                            "sourceRefs": [
                                {"factKey": "organization", "name": "企业名称"}
                            ],
                        }
                    ],
                    "blanksToFill": ["投标报价"],
                },
            )
            self.assertIn("测试企业", result["content"])
            self.assertIn("投标报价", result["blanksToFill"])
            self.assertEqual(result["status"], "generated")

    def test_pi_template_fill_returns_agent_trace(self) -> None:
        client = FakeTemplateFillClient()
        result = PiWritingOrchestrator(client).fill_template_fragment(
            parse_bundle={"parseBundleId": "parse-001"},
            fragment={"templateId": "tpl-001", "editableFields": []},
            section={"sectionId": "c-01", "title": "投标函"},
            project={},
            company={},
            knowledge={},
            workspace_id="w1",
            project_id="p1",
        )
        self.assertEqual(client.calls, ["TableFormFillAgent"])
        self.assertEqual(result["provider"], "pi")
        self.assertEqual(result["phase"], "template_fill")
        self.assertEqual(result["agentResults"][0]["agentId"], "TableFormFillAgent")
        self.assertEqual(result["fieldValues"][0]["value"], "测试企业")

    def test_export_copies_template_table_and_fills_text(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            bundle = self._bundle(root, source)
            fragment = bundle["templateFragments"][0]
            output = assemble_docx(
                {
                    "title": "测试项目",
                    "mode": "open",
                    "plan": [
                        {
                            "sectionId": "c-01",
                            "title": "投标函",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-001",
                            "templateFieldValues": [
                                {
                                    "label": "投标人名称",
                                    "value": "测试企业",
                                    "status": "filled",
                                }
                            ],
                            "content": "备用文本",
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "runtime" / "projects" / "p1" / "outputs",
                root=root,
                template_fragments=[fragment],
            )
            self.assertEqual(output["copiedTemplateCount"], 1)
            exported = Document(output["path"])
            paragraphs = "\n".join(paragraph.text for paragraph in exported.paragraphs)
            tables = "\n".join(
                cell.text
                for table in exported.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("投标函格式", paragraphs)
            self.assertIn("测试企业", paragraphs)
            self.assertIn("固定表头", tables)

    def test_export_uses_manually_edited_template_content_without_duplicate_copy(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            bundle = self._bundle(root, source)
            fragment = bundle["templateFragments"][0]
            output = assemble_docx(
                {
                    "title": "测试项目",
                    "mode": "open",
                    "plan": [
                        {
                            "sectionId": "c-01",
                            "title": "投标函",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-001",
                            "content": "手动修改后的正文",
                            "contentMatchesTemplate": False,
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "runtime" / "projects" / "p1" / "outputs",
                root=root,
                template_fragments=[fragment],
            )
            exported = Document(output["path"])
            paragraphs = "\n".join(paragraph.text for paragraph in exported.paragraphs)
            self.assertIn("手动修改后的正文", paragraphs)
            self.assertNotIn("投标函格式", paragraphs)
            self.assertTrue(any("正文编辑内容" in item for item in output["exportWarnings"]))

    def test_export_preserves_template_format_after_structured_edits(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            source_document = Document(source)
            template_heading = next(
                paragraph for paragraph in source_document.paragraphs
                if paragraph.text == "投标函格式"
            )
            template_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            template_heading.runs[0].bold = True
            template_heading.runs[0].font.size = Pt(16)
            source_document.save(source)
            fragment = self._bundle(root, source)["templateFragments"][0]
            template_blocks = deepcopy(fragment["blocks"])
            template_heading_block = next(
                block for block in template_blocks
                if str(block.get("text") or "") == "投标函格式"
            )
            template_heading_block["text"] = "投标函格式（网页编辑）"
            template_tables = deepcopy(fragment["tables"])
            edited_cell = next(
                cell
                for table in template_tables
                for row in table.get("rows", [])
                for cell in row.get("cells", [])
                if str(cell.get("text") or "") == "________"
            )
            edited_cell["filledText"] = "网页中编辑的表格值"

            output = assemble_docx(
                {
                    "title": "测试项目",
                    "mode": "open",
                    "plan": [
                        {
                            "sectionId": "c-01",
                            "title": "投标函",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-001",
                            "content": "结构化模板编辑",
                            "contentMatchesTemplate": True,
                            "templateBlocks": template_blocks,
                            "templateTables": template_tables,
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "runtime" / "projects" / "p1" / "outputs",
                root=root,
                template_fragments=[fragment],
            )

            exported = Document(output["path"])
            edited_paragraph = next(
                paragraph for paragraph in exported.paragraphs
                if "网页编辑" in paragraph.text
            )
            self.assertEqual(edited_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertTrue(edited_paragraph.runs[0].bold)
            self.assertEqual(edited_paragraph.runs[0].font.size.pt, 16)
            table_text = "\n".join(
                cell.text
                for table in exported.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("固定表头", table_text)
            self.assertIn("网页中编辑的表格值", table_text)
            self.assertEqual(output["copiedTemplateCount"], 1)

    def test_template_preview_keeps_docx_structure_and_editor_bookmarks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            fragment = self._bundle(root, source)["templateFragments"][0]
            preview = render_template_preview_docx(
                fragment,
                {
                    "templateFragmentId": "tpl-001",
                    "templateBlocks": deepcopy(fragment["blocks"]),
                    "templateTables": deepcopy(fragment["tables"]),
                    "templateFieldValues": [],
                    "contentMatchesTemplate": True,
                },
                root=root,
            )
            self.assertIsNotNone(preview)
            assert preview is not None
            preview_document = Document(BytesIO(preview[0]))
            body_xml = preview_document._element.body.xml
            self.assertIn("codex_tpl_p_", body_xml)
            self.assertIn("codex_tpl_t_", body_xml)
            self.assertIn("固定表头", "\n".join(
                cell.text
                for table in preview_document.tables
                for row in table.rows
                for cell in row.cells
            ))

    def test_save_bid_section_persists_structured_template_edits(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = Store(root)
            store.workspaces["w1"] = {
                "id": "w1",
                "bidDocument": {
                    "status": "planned",
                    "plan": [
                        {
                            "sectionId": "c-01",
                            "title": "投标函",
                            "templateFragmentId": "tpl-001",
                            "content": "",
                        }
                    ],
                },
            }

            store.save_bid_section(
                "w1",
                "c-01",
                "投标人名称：修改后的企业",
                template_blocks=[{"bodyChildIndex": 2, "text": "投标人名称：修改后的企业"}],
                template_tables=[],
                content_matches_template=True,
            )

            saved_section = store.workspaces["w1"]["bidDocument"]["plan"][0]
            self.assertTrue(saved_section["contentMatchesTemplate"])
            self.assertEqual(
                saved_section["templateBlocks"][0]["text"],
                "投标人名称：修改后的企业",
            )
            persisted = json.loads((root / "runtime" / "workspaces.json").read_text(encoding="utf-8"))
            persisted_section = persisted["w1"]["bidDocument"]["plan"][0]
            self.assertTrue(persisted_section["contentMatchesTemplate"])

    def test_store_can_fill_template_without_drafting_section(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = self._source_docx(root)
            bundle = self._bundle(root, source)
            bundle.update(
                {
                    "parseBundleId": "parse-001",
                    "schemaVersion": "1.0",
                    "provider": "pi",
                    "status": "completed",
                }
            )
            parse_path = root / "runtime" / "projects" / "p1" / "parse" / "parse-bundle.json"
            parse_path.parent.mkdir(parents=True, exist_ok=True)
            parse_path.write_text(json.dumps(bundle, ensure_ascii=False), encoding="utf-8")

            store = Store(root)
            store.workspaces["w1"] = {
                "id": "w1",
                "projectId": "p1",
                "project": {"id": "p1", "title": "测试项目"},
                "bidDocument": {
                    "status": "planned",
                    "plan": [
                        {
                            "sectionId": "c-01",
                            "title": "投标函",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-001",
                            "content": "",
                            "status": "pending",
                            "contentMatchesTemplate": False,
                        }
                    ],
                },
            }
            store.organization = {"organizationName": "本地企业"}
            store.company = {"company": "本地企业"}
            fake_client = FakeTemplateFillClient()

            self.assertIsNotNone(store.render_bid_template_preview("w1", "c-01"))

            with (
                patch(
                    "backend.app.cloud_agent.config.DifyConfig.from_env",
                    return_value=SimpleNamespace(provider="pi"),
                ),
                patch("backend.app.pi_agent.client.PiClient", return_value=fake_client),
            ):
                result = store.fill_bid_section_template("w1", "c-01")

            assert result is not None
            self.assertTrue(result["filled"])
            self.assertEqual(result["provider"], "pi")
            section = result["section"]
            self.assertEqual(section["templateFillProvider"], "pi")
            self.assertIn("测试企业", section["content"])
            self.assertTrue(section["contentMatchesTemplate"])
            self.assertEqual(fake_client.calls, ["TableFormFillAgent"])

            store.save_bid_section(
                "w1",
                "c-01",
                section["content"],
                template_blocks=section["templateBlocks"],
                template_tables=section["templateTables"],
                content_matches_template=True,
            )
            exported = store.export_bid_document("w1")
            assert exported is not None
            self.assertIn("copiedTemplateCount", exported)
            self.assertIn("exportWarnings", exported)
            self.assertEqual(
                exported["bidDocument"]["copiedTemplateCount"],
                exported["copiedTemplateCount"],
            )
            self.assertFalse(
                any("投标函" in warning and "正文编辑内容" in warning
                    for warning in exported["exportWarnings"])
            )


if __name__ == "__main__":
    unittest.main()
