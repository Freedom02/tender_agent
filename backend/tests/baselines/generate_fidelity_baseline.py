# -*- coding: utf-8 -*-
"""Generate/check deterministic fidelity snapshots for template, table, and evidence flows."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
BASELINE = Path(__file__).resolve().parent / "fidelity_baseline.json"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _digest(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _stable_evidence_path(path: str | Path) -> str:
    """Use a fixture-relative path so snapshots do not contain temp directories."""

    value = str(path or "").replace("\\", "/")
    if not value:
        return ""
    return f"evidence/{Path(value).name}"


def _cell_shape(cell: Any) -> dict[str, Any]:
    span = cell._tc.tcPr.find(qn("w:gridSpan"))
    merge = cell._tc.tcPr.find(qn("w:vMerge"))
    return {
        "text": cell.text,
        "gridSpan": str(span.get(qn("w:val"))) if span is not None else "1",
        "vMerge": str(merge.get(qn("w:val"))) if merge is not None else "",
    }


def _docx_snapshot(raw: bytes | str | Path) -> dict[str, Any]:
    document = Document(BytesIO(raw) if isinstance(raw, bytes) else str(raw))
    paragraphs: list[dict[str, Any]] = []
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("生成时间："):
            continue
        run = paragraph.runs[0] if paragraph.runs else None
        size = run.font.size.pt if run and run.font.size else None
        alignment = paragraph.alignment
        paragraphs.append(
            {
                "text": paragraph.text,
                "alignment": int(alignment) if alignment is not None else None,
                "bold": bool(run.bold) if run else False,
                "fontSize": float(size) if size is not None else None,
            }
        )
    tables = [
        {
            "rows": [
                {"cells": [_cell_shape(cell) for cell in row.cells]}
                for row in table.rows
            ],
            "rowCount": len(table.rows),
            "columnCount": len(table.columns),
        }
        for table in document.tables
    ]
    return {"paragraphs": paragraphs, "tables": tables}


def _stable_image_snapshot(raw: bytes) -> dict[str, Any]:
    """Describe decoded pixels so the baseline ignores PNG encoder details."""

    with Image.open(BytesIO(raw)) as image:
        rgb = image.convert("RGB")
        pixels = rgb.tobytes()
        return {
            "format": str(image.format or ""),
            "mode": str(image.mode),
            "size": [int(image.width), int(image.height)],
            "pixelByteCount": len(pixels),
            "pixelSha256": hashlib.sha256(pixels).hexdigest(),
        }


def _docx_media_snapshot(path: str | Path) -> dict[str, Any]:
    """Capture media structure and decoded pixels for the evidence baseline."""

    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    doc_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    rel_id = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with ZipFile(path) as archive:
        media_names = sorted(
            name for name in archive.namelist() if name.startswith("word/media/")
        )
        relationships = ElementTree.fromstring(
            archive.read("word/_rels/document.xml.rels")
        )
        targets = {
            item.get("Id"): "word/" + str(item.get("Target") or "").lstrip("/")
            for item in relationships.findall(f"{{{rel_ns}}}Relationship")
            if str(item.get("Type") or "").endswith("/image")
        }
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        embedding_order = [
            targets.get(blip.get(f"{{{rel_id}}}embed"), "")
            for blip in document.findall(f".//{{{doc_ns}}}blip")
        ]
        images = [
            {"name": name, **_stable_image_snapshot(archive.read(name))}
            for name in media_names
        ]
    return {
        "mediaNames": media_names,
        "embeddingOrder": embedding_order,
        "images": images,
    }


def _template_snapshot() -> dict[str, Any]:
    from backend.app.bid_writing.assemble import assemble_docx, render_template_preview_docx
    from backend.app.bid_writing.template_reuse import fill_template_fragment
    from backend.tests.fixtures.fidelity_fixtures import template_bundle

    with TemporaryDirectory() as directory:
        root = Path(directory)
        fixture = template_bundle(root)
        fragment = fixture["fragment"]
        fill = fill_template_fragment(
            {"sectionId": "section-fixture", "title": "投标函"},
            fragment,
            {"parsed": {"projectName": "测试项目"}},
            {"organizationName": "测试企业"},
            {},
            agent_result={
                "fieldValues": [
                    {
                        "fieldId": "field-company",
                        "label": "投标人名称",
                        "value": "测试企业",
                        "status": "filled",
                        "sourceRefs": [{"factKey": "organization", "name": "企业名称"}],
                    },
                    {
                        "fieldId": "unknown-field",
                        "label": "未授权字段",
                        "value": "不应填入",
                        "status": "filled",
                        "sourceRefs": [{"factKey": "unknown", "name": "未知"}],
                    },
                ]
            },
        )
        section = {
            "sectionId": "section-fixture",
            "title": "投标函",
            "source": "copied_template",
            "templateFragmentId": "tpl-fixture",
            "content": fill["content"],
            "contentMatchesTemplate": True,
            "templateFieldValues": fill["templateFieldValues"],
            "templateBlocks": fill["templateBlocks"],
            "templateTables": fill["templateTables"],
        }
        output = assemble_docx(
            {
                "title": "测试项目",
                "mode": "open",
                "plan": [section],
                "deviationTables": {"technical": [], "commercial": []},
            },
            root / "output",
            root=root,
            template_fragments=[fragment],
        )
        preview = render_template_preview_docx(fragment, section, root=root)
        assert preview is not None
        preview_bytes, preview_meta = preview
        document = Document(output["path"])
        return {
            "fragment": {
                "copyMode": fragment["copyMode"],
                "bindingStatus": fragment["bindingStatus"],
                "bodyChildStart": fragment["bodyChildStart"],
                "bodyChildEnd": fragment["bodyChildEnd"],
                "fixedText": ["投标函格式", "后续固定内容", "锁定文字"],
                "editableFields": [
                    {key: field.get(key, "") for key in ("fieldId", "label", "sourceKey")}
                    for field in fragment["editableFields"]
                ],
                "tableShape": [
                    {
                        "rowCount": table["rowCount"],
                        "columnCount": table["columnCount"],
                        "gridSpans": [
                            cell.get("gridSpan", "1")
                            for row in table.get("rows", [])
                            for cell in row.get("cells", [])
                        ],
                    }
                    for table in fragment["tables"]
                ],
                "styleAnchor": fragment["blocks"][1]["style"],
            },
            "fill": {
                "content": fill["content"],
                "blanksToFill": fill["blanksToFill"],
                "fieldValues": [
                    {key: field.get(key, "") for key in ("fieldId", "label", "value", "status")}
                    for field in fill["templateFieldValues"]
                ],
                "unknownValueAbsent": "不应填入" not in fill["content"],
            },
            "export": {
                "copiedTemplateCount": output["copiedTemplateCount"],
                "document": _docx_snapshot(output["path"]),
            },
            "preview": {
                "metadata": preview_meta,
                "document": _docx_snapshot(preview_bytes),
            },
        }


def _response_snapshot() -> dict[str, Any]:
    from backend.app.bid_writing.response_table import (
        build_response_table_result,
        build_template_response_table_spec,
        prepare_response_table_inputs,
    )
    from backend.tests.fixtures.fidelity_fixtures import response_table_inputs

    section, fine_review, knowledge = response_table_inputs()
    prepared = prepare_response_table_inputs(section, fine_review, knowledge)
    result = build_response_table_result(
        section,
        prepared,
        {"REQ-001": "支持 SSO 单点登录。"},
    )
    fragment = {
        "tables": [
            {
                "tableId": "TABLE-RESPONSE",
                "rows": [
                    {
                        "rowIndex": 1,
                        "cells": [
                            {"cellIndex": 1, "text": "序号"},
                            {"cellIndex": 2, "text": "技术和服务要求"},
                            {"cellIndex": 3, "text": "投标人响应"},
                            {"cellIndex": 4, "text": "偏离情况"},
                            {"cellIndex": 5, "text": "备注"},
                        ],
                    },
                    {
                        "rowIndex": 2,
                        "cells": [
                            {"cellIndex": 1, "text": ""},
                            {"cellIndex": 2, "text": ""},
                            {"cellIndex": 3, "text": ""},
                            {"cellIndex": 4, "text": ""},
                            {"cellIndex": 5, "text": ""},
                        ],
                    },
                ],
            }
        ]
    }
    mapped = build_template_response_table_spec(fragment, result["requirementResponseRows"])
    markdown_lines = [line for line in result["tableMarkdown"].splitlines() if line.startswith("|")]
    return {
        "preparedRowCount": len(prepared),
        "outputRowCount": len(result["requirementResponseRows"]),
        "markdownColumnCount": len(markdown_lines[0].split("|")[1:-1]),
        "rows": [
            {
                "requirementId": row["requirementId"],
                "matchStatus": row["matchStatus"],
                "needsReview": row["needsReview"],
                "response": row["response"],
                "sourceRefs": [ref.get("evidencePath", "") for ref in row["sourceRefs"]],
            }
            for row in result["requirementResponseRows"]
        ],
        "sourceRefCount": len(result["sourceRefs"]),
        "blanksToFill": result["blanksToFill"],
        "mapped": {
            "tableId": mapped["tableId"] if mapped else None,
            "columns": mapped["columns"] if mapped else None,
            "mappedRowCount": len(mapped["rows"]) if mapped else 0,
            "appendRowCount": len(mapped["appendRows"]) if mapped else 0,
        },
    }


def _evidence_snapshot() -> dict[str, Any]:
    from backend.app.bid_writing.assemble import assemble_docx
    from backend.app.bid_writing.evidence_chapter import build_evidence_chapter_result
    from backend.tests.fixtures.fidelity_fixtures import evidence_fixture

    with TemporaryDirectory() as directory:
        root = Path(directory)
        fixture = evidence_fixture(root)
        chapter = build_evidence_chapter_result(
            {"title": "技术参数证明材料"}, fixture["entries"], fixture["assets"]
        )
        section = {
            "sectionId": "evidence-01",
            "title": "技术参数证明材料",
            "responseType": "attachment",
            "content": chapter["content"],
            "evidenceGroups": chapter["evidenceGroups"],
            "sourceRefs": chapter["sourceRefs"],
        }
        output = assemble_docx(
            {
                "title": "证据测试项目",
                "mode": "open",
                "plan": [section],
                "deviationTables": {"technical": [], "commercial": []},
            },
            root / "output",
            assets=fixture["assets"],
        )
        paragraphs = [paragraph.text for paragraph in Document(output["path"]).paragraphs]
        evidence_markers = [
            paragraph
            for paragraph in paragraphs
            if "产品证明" in paragraph or "证书证明" in paragraph or "待补证明" in paragraph
        ]
        media = _docx_media_snapshot(output["path"])
        expected_images = [
            _stable_image_snapshot(
                Path(fixture["assets"][fixture["refs"][key]]["fixture_page"]).read_bytes()
            )
            for key in ("readable", "second")
        ]
        actual_images = [
            {key: value for key, value in image.items() if key != "name"}
            for image in media["images"]
        ]
        if actual_images != expected_images:
            raise AssertionError("证据页图像素内容未完整进入最终 DOCX 媒体")
        return {
            "files": {
                "readable": bool(fixture["assets"][fixture["refs"]["readable"]]["page_images"]),
                "second": bool(fixture["assets"][fixture["refs"]["second"]]["page_images"]),
                "unreadable": Path(fixture["refs"]["unreadable"]).is_file(),
            },
            "summary": chapter["summary"],
            "insertOrder": [
                _stable_evidence_path(ref.get("evidencePath", ""))
                for group in chapter["evidenceGroups"]
                for ref in group["insertRefs"]
            ],
            "reusedOrder": [
                _stable_evidence_path(ref.get("evidencePath", ""))
                for group in chapter["evidenceGroups"]
                for ref in group["reusedRefs"]
            ],
            "blanksToFill": chapter["blanksToFill"],
            "evidenceMarkers": [
                _stable_evidence_marker(marker, fixture["refs"])
                for marker in evidence_markers
            ],
            "docxMedia": {
                "mediaCount": len(media["mediaNames"]),
                "mediaNames": media["mediaNames"],
                "embeddingOrder": media["embeddingOrder"],
                "images": media["images"],
                "expectedImages": expected_images,
                "missingHasNoPageImages": not bool(
                    fixture["assets"][fixture["refs"]["unreadable"]]["page_images"]
                ),
            },
        }


def _stable_evidence_marker(marker: str, refs: dict[str, str]) -> str:
    """Normalize generated evidence text without changing its assertions."""

    normalized = str(marker)
    for path in refs.values():
        normalized = normalized.replace(
            str(path), _stable_evidence_path(path)
        )
        normalized = normalized.replace(
            str(path).replace("\\", "/"), _stable_evidence_path(path)
        )
    return normalized


def snapshots() -> dict[str, Any]:
    result = {
        "template": _template_snapshot(),
        "responseTable": _response_snapshot(),
        "evidence": _evidence_snapshot(),
    }
    return {
        "schemaVersion": "1.0",
        "snapshots": result,
        "sha256": {name: _digest(value) for name, value in result.items()},
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="只比对，不改写基线")
    args = parser.parse_args()
    current = snapshots()
    if args.check:
        if not BASELINE.is_file():
            print(f"缺少基线文件：{BASELINE}")
            return 1
        expected = json.loads(BASELINE.read_text(encoding="utf-8"))
        if current != expected:
            print("保真基线不一致：")
            print(json.dumps({"expected": expected, "actual": current}, ensure_ascii=False, indent=2))
            return 1
        print("三组保真基线全部一致（字段快照 + sha256）。")
        return 0
    BASELINE.write_text(
        json.dumps(current, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"已写入保真基线：{BASELINE}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
