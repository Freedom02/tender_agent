from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from pydantic import ValidationError

from backend.app.agent_runtime.contracts import build_parse_bundle
from backend.app.agent_runtime.parse_artifact import finalize_native_parse_bundle
from backend.app.agent_runtime.parse_semantics import (
    ClauseExtraction, EvaluationExtraction, SemanticReview, TableBinding,
    build_source_batches, hydrate_source_traces, namespace_clause_batch, validate_clause_batch,
)
from backend.app.agent_runtime.resolved_tender import RESOLVED_TENDER_V2, ResolvedTenderV2
from backend.app.bid_writing.resolved_input import resolve_tender
from backend.app.bid_writing.template_fragments import attach_template_fragments_to_bundle
from backend.app.cloud_agent.normalizers import normalize_parse_result
from backend.app.cloud_agent.run_repository import RunRepository
from backend.app.cloud_agent.run_service import CloudRunService
from backend.app.pi_agent.agents import PiAgentMethods
from backend.app.pi_agent.client import PiClient
from backend.app.pi_agent.config import PiConfig
from backend.app.pi_agent.orchestrator import PiParseOrchestrationError
from backend.app.pi_agent.resolved_parse import PiResolvedParseOrchestrator, compatibility_projection
from backend.app.pi_agent.transport import PolicyBlockedError, agent_id_from_prompt
from backend.app.store import Store
from backend.app.tender_matching.document_parser import parse_docx
from backend.tests.fixtures.resolved_tender_fixtures import digital_twin_v2


def bare_traces(value):
    if isinstance(value, list):
        return [bare_traces(item) for item in value]
    if isinstance(value, dict):
        return {key: [{"segmentId": trace["segmentId"], "quote": trace["quote"]} for trace in item]
                if key in {"sourceTrace", "evaluationMethodSourceTrace"} else bare_traces(item)
                for key, item in value.items()}
    return value


class AnnotatedProtocolClient(PiAgentMethods):
    """Replay manual annotations to test protocol wiring, never model accuracy."""

    def __init__(self, *, failures=(), review_issue=False):
        self.sample = digital_twin_v2()
        self.config = SimpleNamespace(model="offline-protocol")
        self.failures = set(failures)
        self.review_issue = review_issue
        self.calls = []
        self.inputs = []

    def record(self, name):
        self.calls.append(name)
        if name in self.failures:
            raise PolicyBlockedError("offline rejection")

    def _parse_file(self, path):
        blocks = deepcopy(self.sample["documentBlocks"])
        blocks["text"] = "\n".join(unit["text"] for unit in blocks["units"])
        return blocks

    def run_json(self, prompt, **kwargs):
        self.record("structure")
        return {"documentStructure": {"sections": []}, "projectMeta": self.sample["projectMeta"]}

    def detect_templates(self, **kwargs):
        self.record("templates")
        return {"hasExplicitFormat": True, "requiredOutline": [], "requiredTables": [],
                "templates": [], "submissionStructure": deepcopy(self.sample["submissionStructure"])}

    def extract_tender_clauses(self, *, inputs, correction=""):
        self.record("clauses")
        self.inputs.append(deepcopy(inputs))
        requirements = [deepcopy(item) for item in self.sample["requirements"] if any(
            trace["segmentId"] == view["segmentId"] and trace["quote"] in view["text"]
            for trace in item["sourceTrace"] for view in inputs["views"])]
        requirement_ids = {item["requirementId"] for item in requirements}
        evidence = [item for item in self.sample["evidenceRequirements"] if item["requirementId"] in requirement_ids]
        assessments = []
        for view in inputs["views"]:
            matched = [item["requirementId"] for item in requirements if any(
                trace["segmentId"] == view["segmentId"] and trace["quote"] in view["text"] for trace in item["sourceTrace"])]
            assessments.append({"viewId": view["viewId"], "disposition": "extracted" if matched else "context",
                                "requirementIds": matched, "reason": "人工标注协议回放"})
        return bare_traces({"requirements": requirements, "evidenceRequirements": evidence, "assessments": assessments})

    def extract_tender_evaluation(self, *, inputs, correction=""):
        self.record("evaluation")
        mapping = {item["requirementId"].split(":", 1)[-1]: item["requirementId"] for item in inputs["requirements"]}
        rules = [deepcopy(item) for item in self.sample["evaluationRules"] if all(item_id in mapping for item_id in item["requirementIds"])]
        for rule in rules:
            rule["requirementIds"] = [mapping[item_id] for item_id in rule["requirementIds"]]
        method = any(view["segmentId"] == "SEG-000510" for view in inputs["views"])
        return bare_traces({"evaluationMethod": "lowest_evaluated_price" if method else "unknown",
                            "evaluationMethodSourceTrace": self.sample["evaluationMethodSourceTrace"] if method else [],
                            "evaluationRules": rules})

    def repair_tender_clauses(self, *, inputs, correction=""):
        corrected = self.extract_tender_clauses(inputs=inputs, correction=correction)
        fields = {"requirements": ("requirementReplacements", "requirementId"),
                  "evidenceRequirements": ("evidenceReplacements", "evidenceRequirementId"),
                  "assessments": ("assessmentReplacements", "viewId")}
        repair = {"addedAssessments": [item for item in corrected["assessments"] if item["viewId"] in inputs["missingViewIds"]]}
        for target in inputs["targets"]:
            field, identifier = fields[target["collection"]]
            value = next(item for item in corrected[target["collection"]] if item[identifier] == target["value"][identifier])
            repair.setdefault(field, []).append({"targetId": target["targetId"], "value": value})
        return repair

    def bind_tender_response_tables(self, *, inputs, correction=""):
        self.record("tables")
        mapping = {item["requirementId"].split(":", 1)[-1]: item["requirementId"] for item in inputs["requirements"]}
        tables = deepcopy(self.sample["responseTableSchemas"])
        for table in tables:
            table["columns"] = [{**column, "key": f"column-{position}"} for position, column in enumerate(table["columns"], 1)]
            table["rowBindings"] = [{**row, "requirementId": mapping[row["requirementId"]]} for row in table["rowBindings"] if row["requirementId"] in mapping]
        return bare_traces({"responseTableSchemas": [table for table in tables if table["rowBindings"]]})

    def review_tender_semantics(self, *, inputs, correction=""):
        self.record("review")
        return {"status": "review" if self.review_issue else "pass",
                "reviewedViewIds": [view["viewId"] for view in inputs["views"]],
                "reviewedRequirementIds": [item["requirementId"] for item in inputs["requirements"]],
                "findings": [{"code": "manual_semantic_issue", "severity": "review", "path": inputs["batchId"],
                              "message": "独立复核发现遗漏"}] if self.review_issue else []}

    def fragments(self, bundle, **kwargs):
        return {**bundle, "templateFragments": deepcopy(self.sample["templateFragments"])}


class NativeTransportIntegrationTests(unittest.TestCase):
    def client(self):
        return PiClient(PiConfig(enabled=True, command="pi", provider="", model="offline-protocol",
                                 api_key="", timeout_seconds=180, working_dir=""))

    def test_real_client_transport_preserves_strict_node_payloads(self):
        client = self.client()
        cases = (
            (client.extract_tender_clauses, ClauseExtraction,
             {"requirements": [], "evidenceRequirements": [], "assessments": []}),
            (client.extract_tender_evaluation, EvaluationExtraction, {"evaluationRules": []}),
            (client.bind_tender_response_tables, TableBinding, {"responseTableSchemas": []}),
            (client.review_tender_semantics, SemanticReview,
             {"status": "pass", "reviewedViewIds": [], "reviewedRequirementIds": [], "findings": []}),
        )
        with patch.dict("os.environ", {"PI_AGENT_VALIDATION": "1"}):
            for method, model, payload in cases:
                with self.subTest(method=method.__name__), patch.object(client, "_run_json_raw", return_value=deepcopy(payload)) as raw:
                    result = method(inputs={"views": []})
                    self.assertEqual(result, payload)
                    model.model_validate(result)
                    self.assertEqual(raw.call_args.kwargs["tools"], "")
                    self.assertEqual(raw.call_args.kwargs["extensions"], ())
                    self.assertEqual(raw.call_args.kwargs["timeout_seconds"], 180)
                with self.subTest(unexpected_field=method.__name__), patch.object(client, "_run_json_raw", return_value={**payload, "unexpected": True}):
                    with self.assertRaises(ValidationError):
                        model.model_validate(method(inputs={"views": []}))

    def test_default_parse_entry_uses_real_transport_for_all_six_node_types(self):
        client = self.client()
        replay = AnnotatedProtocolClient()
        agents = []
        methods = {
            "TenderClauseAgent": replay.extract_tender_clauses,
            "TenderEvaluationAgent": replay.extract_tender_evaluation,
            "TenderTableBindingAgent": replay.bind_tender_response_tables,
            "TenderSemanticReviewAgent": replay.review_tender_semantics,
        }

        def raw_output(prompt, **kwargs):
            agent_id = agent_id_from_prompt(prompt)
            agents.append(agent_id)
            if agent_id == "DocumentStructureAgent":
                return replay.run_json(prompt)
            if agent_id == "TemplateDetectAgent":
                return replay.detect_templates()
            inputs = json.loads(prompt.split("【本次输入】\n", 1)[1].split("\n【上次校验反馈", 1)[0])
            return methods[agent_id](inputs=inputs)

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.txt"
            path.write_text("离线真实客户端接线验证", encoding="utf-8")
            with patch.dict("os.environ", {"PI_AGENT_VALIDATION": "1"}), \
                 patch.object(client, "_parse_file", side_effect=replay._parse_file), \
                 patch.object(client, "_run_json_raw", side_effect=raw_output), \
                 patch("backend.app.pi_agent.resolved_parse.attach_template_fragments_to_bundle", side_effect=replay.fragments):
                result = client.parse_tender(path, workspace_id="w1", project_id="p1")
        self.assertEqual(set(agents), {"DocumentStructureAgent", "TemplateDetectAgent", *methods})
        self.assertEqual(len(agents), 6)
        self.assertTrue(result["parseReviewReport"]["readyForWriting"], result["parseReviewReport"])
        self.assertEqual(len(result["resolvedTenderV2"]["requirements"]), 10)
        self.assertEqual(len(result["resolvedTenderV2"]["responseTableSchemas"]), 2)


class SourceBatchTests(unittest.TestCase):
    def test_full_text_and_late_requirements_survive_batching(self):
        units = [{"id": f"block-{position:04d}", "order": position, "text": f"第{position}条 " + "完整正文" * 50} for position in range(150)]
        batches = build_source_batches({"units": units}, max_chars=4000)
        views = [view for batch in batches for view in batch["views"]]
        self.assertGreater(len(batches), 6)
        self.assertEqual(len(views), 150)
        self.assertEqual([view["text"] for view in views], [unit["text"] for unit in units])
        self.assertIn("第149条", views[-1]["text"])
        self.assertTrue(batches[1]["contextViews"])
        self.assertTrue(all(len(json.dumps(batch["views"], ensure_ascii=False)) < 4000 for batch in batches))

    def test_oversized_table_row_is_split_without_losing_characters(self):
        text = "开头\n" + "参数10万、帧率60fps。" * 1000 + "\n结尾"
        blocks = {"units": [{"id": "row-1", "contentType": "table_row", "tableId": "table-1", "rowIndex": 1, "text": text},
                            {"id": "cell-1", "contentType": "table_cell", "tableId": "table-1", "rowIndex": 1, "text": text}]}
        views = [view for batch in build_source_batches(blocks, max_chars=4000) for view in batch["views"]]
        self.assertGreater(len(views), 1)
        self.assertEqual("".join(view["text"] for view in views), text)
        self.assertTrue(all(view["segmentId"] == "row-1" for view in views))

    def test_cell_only_sources_are_retained(self):
        blocks = {"tables": [{"id": "table-1", "bodyChildIndex": 4, "rows": [{"rowIndex": 0, "cells": [
            {"id": "cell-1", "cellIndex": 0, "text": "原列名"}]}]}]}
        view = build_source_batches(blocks)[0]["views"][0]
        self.assertEqual((view["segmentId"], view["bodyChildIndex"], view["tableHeaders"]), ("cell-1", 4, ["原列名"]))

    def test_hydration_adds_real_identity_and_rejects_forgery(self):
        blocks = {"units": [{"id": "p1", "text": "真实参数60fps", "paragraphIndex": 3}, {"id": "p2", "text": "另一个条款"}]}
        source = {"fileId": "file1", "sha256": "hash1"}
        view = {"segmentId": "p1", "text": "真实参数60fps"}
        payload = {"sourceTrace": [{"segmentId": "p1", "quote": "60fps"}]}
        result = hydrate_source_traces(payload, document_blocks=blocks, source_file=source, views=[view])
        self.assertEqual(result["sourceTrace"][0]["documentVersionId"], "hash1")
        self.assertEqual(result["sourceTrace"][0]["paragraphIndex"], 3)
        self.assertNotIn("fileId", payload["sourceTrace"][0])
        for trace in ({"segmentId": "absent", "quote": "60fps"}, {"segmentId": "p1", "quote": "30fps"},
                      {"segmentId": "p2", "quote": "另一个条款"}, {"segmentId": "p1", "quote": "60fps", "fileId": "fake"},
                      {"segmentId": "p1", "quote": "60fps", "paragraphIndex": 7}):
            with self.subTest(trace=trace), self.assertRaises(ValueError):
                hydrate_source_traces({"sourceTrace": [trace]}, document_blocks=blocks, source_file=source, views=[view])

    def test_partial_view_cannot_cite_other_part_of_same_block(self):
        with self.assertRaisesRegex(ValueError, "outside_view"):
            hydrate_source_traces({"sourceTrace": [{"segmentId": "p1", "quote": "末尾"}]},
                                 document_blocks={"units": [{"id": "p1", "text": "开头及末尾"}]},
                                 source_file={}, views=[{"segmentId": "p1", "text": "开头"}])

    def test_assessment_coverage_and_relationships_are_enforced(self):
        client = AnnotatedProtocolClient()
        batch = build_source_batches(client.sample["documentBlocks"], max_chars=90000)[0]
        payload = client.extract_tender_clauses(inputs=batch)
        clauses = validate_clause_batch(payload, batch)
        namespaced = namespace_clause_batch(clauses, "batch-2")
        self.assertTrue(all(item.requirementId.startswith("batch-2:") for item in namespaced.requirements))
        self.assertEqual(namespaced.evidenceRequirements[0].alternativeGroupId, "batch-2:entity-registration")
        ResolvedTenderV2(parseBundleId="test", requirements=namespaced.requirements, evidenceRequirements=namespaced.evidenceRequirements)
        for mutation in ("missing", "duplicate", "unknown_ref", "no_link", "wrong_source"):
            invalid = deepcopy(payload)
            if mutation == "missing":
                invalid["assessments"].pop()
            elif mutation == "duplicate":
                invalid["assessments"].append(deepcopy(invalid["assessments"][0]))
            else:
                assessment = next(item for item in invalid["assessments"] if item["requirementIds"])
                if mutation == "unknown_ref":
                    assessment["requirementIds"] = ["absent"]
                elif mutation == "no_link":
                    assessment.update(disposition="context", requirementIds=[])
                else:
                    assessment["requirementIds"] = ["technical-0014"]
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                validate_clause_batch(invalid, batch)


class NativeOrchestrationTests(unittest.TestCase):
    def parse(self, client=None, *, budget=90000):
        client = client or AnnotatedProtocolClient()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.txt"
            path.write_text("协议测试原文件", encoding="utf-8")
            with patch("backend.app.pi_agent.resolved_parse.attach_template_fragments_to_bundle", side_effect=client.fragments):
                events = []
                result = PiResolvedParseOrchestrator(client, batch_chars=budget).parse_tender(
                    path, workspace_id="w1", project_id="p1", project_context={"sourceFileId": "source-1"},
                    progress_callback=lambda *event: events.append(event))
        return result, client, events

    def test_native_default_entry_preserves_manual_semantics_in_protocol_replay(self):
        result, client, events = self.parse()
        tender = ResolvedTenderV2.model_validate(result["resolvedTenderV2"])
        self.assertEqual(tender.origin, "native")
        self.assertEqual(len(tender.requirements), 10)
        self.assertEqual(len(tender.evidenceRequirements), 6)
        self.assertEqual(len(tender.responseTableSchemas), 2)
        self.assertEqual(tender.evaluationMethod, "lowest_evaluated_price")
        self.assertEqual(result["evaluationItems"], [])
        self.assertTrue(result["parseReviewReport"]["readyForWriting"], result["parseReviewReport"])
        screenshot = next(item for item in tender.evidenceRequirements if item.type == "screenshot")
        self.assertEqual(screenshot.minimumQuantity, 4)
        self.assertIn("加盖投标人公章", screenshot.verificationRules)
        rendering = next(item for item in tender.requirements if item.requirementId.endswith("technical-0014"))
        self.assertEqual([(parameter.value, parameter.unit) for parameter in rendering.parameters], [("100000", "个"), ("60", "fps")])
        self.assertEqual(len(result["requirements"]), 6)
        self.assertTrue(all(item["responseStage"] == "bid_submission" for item in result["requirements"]))
        normalized = normalize_parse_result(result).model_dump(mode="json")
        self.assertTrue(all(isinstance(item["mandatory"], bool) for item in normalized["requirements"]))
        self.assertEqual(events[-1][2:], (6, 6))
        self.assertEqual(set(client.calls), {"structure", "clauses", "evaluation", "templates", "tables", "review"})

    def test_multiple_batches_merge_references_tables_and_progress(self):
        result, client, events = self.parse(budget=5000)
        coverage = result["parseReviewReport"]["coverage"]
        self.assertGreater(coverage["batchCount"], 1)
        self.assertEqual(coverage["sourceViewCount"], coverage["assessedViewCount"])
        tender = ResolvedTenderV2.model_validate(result["resolvedTenderV2"])
        self.assertEqual(len(tender.requirements), 10)
        self.assertEqual(len(tender.responseTableSchemas), 2)
        self.assertEqual(events[-1][2], events[-1][3])
        self.assertGreater(events[-1][3], 6)
        self.assertEqual([event[2] for event in events], sorted(event[2] for event in events))

    def test_independent_review_findings_prevent_readiness(self):
        result, _, _ = self.parse(AnnotatedProtocolClient(review_issue=True))
        self.assertFalse(result["parseReviewReport"]["readyForWriting"])
        self.assertEqual(result["parseReviewReport"]["status"], "review_required")
        self.assertIn("manual_semantic_issue", {issue["code"] for issue in result["parseReviewReport"]["issues"]})

    def test_incomplete_semantic_review_is_retried_then_recorded_as_failure(self):
        client = AnnotatedProtocolClient()
        client.review_tender_semantics = lambda **kwargs: {"status": "pass", "reviewedViewIds": [], "reviewedRequirementIds": [], "findings": []}
        result, _, _ = self.parse(client)
        record = next(item for item in result["agentResults"] if item["agentId"] == "TenderSemanticReviewAgent")
        self.assertEqual(record["status"], "failed")
        self.assertEqual(len(record["output"]["attempts"]), 2)
        self.assertFalse(result["parseReviewReport"]["readyForWriting"])

    def test_rejected_extraction_keeps_unknown_source_without_false_duties(self):
        result, client, _ = self.parse(AnnotatedProtocolClient(failures={"clauses"}))
        self.assertEqual(client.calls.count("clauses"), 1)
        self.assertEqual(result["requirements"], [])
        self.assertTrue(result["resolvedTenderV2"]["requirements"])
        self.assertTrue(all(item["actor"] == "unknown" and item["responseRequired"] is None for item in result["resolvedTenderV2"]["requirements"]))
        self.assertFalse(result["parseReviewReport"]["readyForWriting"])
        self.assertGreater(result["parseReviewReport"]["coverage"]["unresolvedViewCount"], 0)

    def test_all_nodes_failed_raises_instead_of_reporting_success(self):
        with self.assertRaises(PiParseOrchestrationError):
            self.parse(AnnotatedProtocolClient(failures={"structure", "clauses", "evaluation", "templates", "review"}))

    def test_bad_quote_is_corrected_once_and_both_attempts_are_preserved(self):
        client = AnnotatedProtocolClient()
        original = client.extract_tender_clauses
        corrections = []

        def extract(**kwargs):
            payload = original(**kwargs)
            corrections.append(kwargs["correction"])
            if not kwargs["correction"]:
                payload["requirements"][0]["sourceTrace"][0]["quote"] = "虚构原文"
            return payload

        client.extract_tender_clauses = extract
        result, _, _ = self.parse(client)
        record = next(item for item in result["agentResults"] if item["agentId"] == "TenderClauseAgent")
        self.assertEqual(len(record["output"]["attempts"]), 2)
        self.assertIn("outside_view", corrections[1])
        self.assertEqual(record["status"], "success")

    def test_source_mutation_during_parsing_is_rejected(self):
        client = AnnotatedProtocolClient()
        original = client._parse_file

        def parse(path):
            result = original(path)
            path.write_text("changed", encoding="utf-8")
            return result

        client._parse_file = parse
        with self.assertRaisesRegex(PiParseOrchestrationError, "发生变化"):
            self.parse(client)

    def test_time_value_cannot_become_score_in_native_extraction(self):
        client = AnnotatedProtocolClient()
        original = client.extract_tender_evaluation

        def evaluation(**kwargs):
            result = original(**kwargs)
            result["evaluationMethod"] = "unknown"
            result["evaluationMethodSourceTrace"] = []
            result["evaluationRules"][0]["ruleType"] = "subjective_scoring"
            result["evaluationRules"][0]["maxScore"] = {"value": 30.0, "sourceTrace": result["evaluationRules"][0]["sourceTrace"]}
            return result

        client.extract_tender_evaluation = evaluation
        result, _, _ = self.parse(client)
        record = next(item for item in result["agentResults"] if item["agentId"] == "TenderEvaluationAgent")
        self.assertEqual(record["status"], "failed")
        self.assertIn("score_value_unverified", record["output"]["attempts"][-1]["error"])
        self.assertEqual(result["evaluationItems"], [])

    def test_source_requirement_table_cannot_replace_original_response_template(self):
        client = AnnotatedProtocolClient()
        original = client.bind_tender_response_tables

        def bind(**kwargs):
            result = original(**kwargs)
            result["responseTableSchemas"][0]["sourceTableId"] = "TABLE-018"
            return result

        client.bind_tender_response_tables = bind
        result, _, _ = self.parse(client)
        record = next(item for item in result["agentResults"] if item["agentId"] == "TenderTableBindingAgent")
        self.assertEqual(record["status"], "failed")
        self.assertEqual(result["resolvedTenderV2"]["responseTableSchemas"], [])
        self.assertFalse(result["parseReviewReport"]["readyForWriting"])

    def test_source_mutation_during_model_calls_is_rejected(self):
        client = AnnotatedProtocolClient()
        original = client._parse_file
        source_path = []

        def parse(path):
            source_path.append(path)
            return original(path)

        original_review = client.review_tender_semantics

        def review(**kwargs):
            source_path[0].write_text("changed while reviewing", encoding="utf-8")
            return original_review(**kwargs)

        client._parse_file = parse
        client.review_tender_semantics = review
        with self.assertRaisesRegex(PiParseOrchestrationError, "发生变化"):
            self.parse(client)

    def test_agent_run_ids_are_scoped_to_unique_parse_bundles(self):
        first, _, _ = self.parse()
        second, _, _ = self.parse()
        self.assertTrue(all(item["runId"].startswith(first["parseBundleId"]) for item in first["agentResults"]))
        self.assertFalse({item["runId"] for item in first["agentResults"]} & {item["runId"] for item in second["agentResults"]})

    def test_public_client_entry_uses_native_orchestrator(self):
        client = AnnotatedProtocolClient()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.txt"
            path.write_text("test", encoding="utf-8")
            with patch("backend.app.pi_agent.resolved_parse.attach_template_fragments_to_bundle", side_effect=client.fragments):
                result = client.parse_tender(path, workspace_id="w1", project_id="p1")
        self.assertEqual(result["resolvedTenderV2"]["schemaVersion"], RESOLVED_TENDER_V2)


class NativePersistenceTests(unittest.TestCase):
    def bundle(self):
        tender = digital_twin_v2()
        tender.update(projectId="p1", workspaceId="w1")
        result = {**compatibility_projection(ResolvedTenderV2.model_validate(tender)),
                  "parseBundleId": tender["parseBundleId"], "provider": "pi", "sourceFile": tender["sourceFile"],
                  "documentBlocks": tender["documentBlocks"], "submissionStructure": tender["submissionStructure"],
                  "templateFragments": tender["templateFragments"], "resolvedTenderV2": tender}
        return build_parse_bundle(result=normalize_parse_result(result).model_dump(mode="json"), project_id="p1", workspace_id="w1")

    def test_native_saved_with_sidecars_and_resolved_without_legacy_migration(self):
        bundle = self.bundle()
        bundle["templateFragments"][0]["rawXml"] = "<w:p>原模板锁定文字</w:p>"
        bundle["templateFragmentsArtifact"] = "runtime/templates/index.json"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = Store(root)
            saved = store.save_parse_bundle(bundle)
            native = json.loads((root / saved["resolvedTenderV2Artifact"]).read_text(encoding="utf-8"))
            review = json.loads((root / saved["resolvedTenderV2ReviewArtifact"]).read_text(encoding="utf-8"))
            self.assertEqual(native, saved["resolvedTenderV2"])
            self.assertTrue(review["readyForWriting"], review)
            self.assertEqual(native["templateFragments"][0]["rawXml"], bundle["templateFragments"][0]["rawXml"])
            self.assertEqual(native["templateFragmentsArtifact"], "runtime/templates/index.json")
            self.assertEqual({item["type"] for item in saved["artifacts"]}, {"parse_bundle", "resolved_tender_v2", "resolved_tender_v2_review"})
            resolved = resolve_tender({}, {}, parse_bundle=saved, schema_version=RESOLVED_TENDER_V2,
                                      company={"materials": [{"path": "real-proof.pdf"}]})
            self.assertEqual(resolved["requirements"], native["requirements"])
            self.assertEqual(resolved["origin"], "native")
            self.assertEqual(resolved["company"]["materials"][0]["path"], "real-proof.pdf")

    def test_invalid_embedded_native_never_falls_back_to_legacy(self):
        for invalid in (None, {}, {"schemaVersion": "wrong"}):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                resolve_tender({}, {}, parse_bundle={"resolvedTenderV2": invalid}, schema_version=RESOLVED_TENDER_V2)

    def test_envelope_identity_and_source_mismatches_are_rejected(self):
        for field in ("parseBundleId", "projectId", "workspaceId", "sha256", "fileId"):
            bundle = self.bundle()
            if field in {"sha256", "fileId"}:
                bundle["sourceFile"][field] = "wrong"
            else:
                bundle[field] = "wrong"
            with self.subTest(field=field), self.assertRaises(ValueError):
                finalize_native_parse_bundle(bundle)

    def test_unassigned_project_can_be_bound_once(self):
        bundle = self.bundle()
        bundle["resolvedTenderV2"]["projectId"] = ""
        bundle["resolvedTenderV2"]["workspaceId"] = ""
        finalized = finalize_native_parse_bundle(bundle)
        self.assertEqual(finalized["resolvedTenderV2"]["projectId"], "p1")
        self.assertEqual(finalized["resolvedTenderV2"]["workspaceId"], "w1")

    def test_real_docx_template_assets_survive_final_native_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "source.docx"
            document = Document()
            document.add_heading("第一章 投标文件格式", level=1)
            document.add_paragraph("投标函格式")
            document.add_paragraph("投标人名称：________")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "固定表头"
            table.cell(0, 1).text = "填写项"
            table.cell(1, 0).text = "项目名称"
            table.cell(1, 1).text = "________"
            document.save(path)
            blocks = parse_docx(path.read_bytes())
            tender = ResolvedTenderV2(parseBundleId="parse-real", projectId="p1", workspaceId="w1", documentBlocks=blocks,
                                      sourceFile={"fileId": "file-1", "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
            bundle = build_parse_bundle(result={"parseBundleId": tender.parseBundleId, "documentBlocks": blocks,
                "sourceFile": tender.sourceFile, "resolvedTenderV2": tender.model_dump(mode="json"),
                "bidFormatSpec": {"templates": [{"templateId": "tpl-1", "title": "投标函格式", "sourceLocator": "投标函格式", "templateType": "fixed_template"}]}},
                project_id="p1", workspace_id="w1")
            bundle = attach_template_fragments_to_bundle(bundle, root=root, source_path=path)
            saved = Store(root).save_parse_bundle(bundle)
            fragment = saved["resolvedTenderV2"]["templateFragments"][0]
            self.assertEqual(fragment["copyMode"], "ooxml_body_range")
            self.assertTrue((root / fragment["fragmentArtifact"]).is_file())
            self.assertTrue((root / saved["resolvedTenderV2"]["templateFragmentsArtifact"]).is_file())
            self.assertTrue(fragment["tables"])
            self.assertEqual(saved["resolvedTenderV2"]["sourceFile"]["storagePath"], "source.docx")

    def test_unknown_compatibility_obligation_is_not_promoted(self):
        tender = ResolvedTenderV2(parseBundleId="test", requirements=[{"requirementId": "r1", "text": "未确认要求"}])
        self.assertEqual(compatibility_projection(tender)["requirements"], [])


class NativeCloudIntegrationTests(unittest.TestCase):
    def test_service_updates_variable_progress_and_native_template_snapshot(self):
        client = AnnotatedProtocolClient()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            repository = RunRepository(root / "runtime")
            updates = []
            original = repository.update

            def update(run_id, **values):
                updates.append(deepcopy(values))
                return original(run_id, **values)

            repository.update = update
            service = CloudRunService(root, repository=repository, pi_client=client, agent_provider="pi")
            run = service.create_run(workspace_id="w1", project_id="p1", filename="source.txt", content=b"offline-test")
            with patch("backend.app.pi_agent.resolved_parse.attach_template_fragments_to_bundle", side_effect=client.fragments):
                finished = service.process_run(run["runId"])
            self.assertEqual(finished["status"], "outline_ready", finished.get("error"))
            payload = json.loads((root / finished["parseBundleArtifact"]).read_text(encoding="utf-8"))
            self.assertEqual(payload["resolvedTenderV2"]["templateFragments"], payload["templateFragments"])
            self.assertEqual(finished["result"]["resolvedTenderV2"], payload["resolvedTenderV2"])
            progress = [item["progress"] for item in updates if "progress" in item]
            self.assertEqual(progress, sorted(progress))
            self.assertIn(96, progress)
            self.assertEqual(progress[-1], 100)


if __name__ == "__main__":
    unittest.main()
