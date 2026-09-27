from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.app.bid_writing.assemble import assemble_docx
from backend.app.bid_writing.response_table import build_template_response_table_spec


class ResponseTableDocxExportTests(unittest.TestCase):
    def test_existing_template_rows_are_filled_and_missing_rows_are_cloned(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_path = root / "tender.docx"
            source = Document()
            source.add_paragraph("技术和服务要求响应表")
            table = source.add_table(rows=3, cols=5)
            values = [
                ["序号", "招标要求", "投标应答", "偏离情况", "备注"],
                ["1", "平台应支持三维可视化", "", "", ""],
                ["2", "提供现场培训服务", "", "", ""],
            ]
            for row_index, row_values in enumerate(values):
                for column_index, value in enumerate(row_values):
                    table.cell(row_index, column_index).text = value
            source.save(source_path)

            fragment = {
                "templateId": "tpl-response",
                "fragmentType": "fixed_table",
                "copyMode": "ooxml_body_range",
                "sourceFilePath": str(source_path),
                "bodyChildStart": 1,
                "bodyChildEnd": 1,
                "tables": [
                    {
                        "tableId": "TABLE-001",
                        "rows": [
                            {
                                "rowIndex": row_index,
                                "cells": [
                                    {"cellIndex": column_index, "text": value}
                                    for column_index, value in enumerate(row_values, start=1)
                                ],
                            }
                            for row_index, row_values in enumerate(values, start=1)
                        ],
                    }
                ],
            }
            response_rows = [
                {
                    "sequence": 1,
                    "requirementId": "req-1",
                    "requirement": "平台应支持三维可视化",
                    "response": "我方平台支持三维可视化。",
                    "deviation": "无偏离",
                    "remark": "产品参数表",
                },
                {
                    "sequence": 2,
                    "requirementId": "req-2",
                    "requirement": "提供现场培训服务",
                    "response": "我方提供现场培训服务。",
                    "deviation": "待补充材料",
                    "remark": "待补充培训证明",
                },
                {
                    "sequence": 3,
                    "requirementId": "req-3",
                    "requirement": "平台应支持数据接入",
                    "response": "我方平台支持数据接入。",
                    "deviation": "无偏离",
                    "remark": "产品参数表",
                },
            ]
            table_spec = build_template_response_table_spec(fragment, response_rows)
            self.assertIsNotNone(table_spec)
            assert table_spec is not None
            self.assertEqual(
                [item["rowIndex"] for item in table_spec["rows"]],
                [2, 3],
            )

            output = assemble_docx(
                {
                    "title": "示例项目",
                    "plan": [
                        {
                            "sectionId": "tech-response",
                            "title": "技术和服务要求响应表",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-response",
                            "templateResponseTable": table_spec,
                            "requirementResponseRows": response_rows,
                            "content": "备用正文",
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "outputs",
                root=root,
                template_fragments=[fragment],
            )

            exported = Document(output["path"])
            response_table = next(
                table
                for table in exported.tables
                if any(cell.text == "投标应答" for row in table.rows for cell in row.cells)
            )
            exported_rows = [
                [cell.text for cell in row.cells]
                for row in response_table.rows
            ]
            self.assertEqual(len(exported_rows), 4)
            self.assertEqual(exported_rows[1][2], "我方平台支持三维可视化。")
            self.assertEqual(
                exported_rows[2][2],
                "我方提供现场培训服务。",
                repr(exported_rows),
            )
            self.assertEqual(exported_rows[3][1], "平台应支持数据接入")
            self.assertEqual(exported_rows[3][2], "我方平台支持数据接入。")

    def test_unrecognized_template_table_appends_standard_response_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_path = root / "tender.docx"
            source = Document()
            source.add_paragraph("技术和服务要求响应表")
            table = source.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "固定表头"
            table.cell(0, 1).text = "填写项"
            source.save(source_path)

            fragment = {
                "templateId": "tpl-unknown",
                "fragmentType": "fixed_table",
                "copyMode": "ooxml_body_range",
                "sourceFilePath": str(source_path),
                "bodyChildStart": 1,
                "bodyChildEnd": 1,
                "tables": [
                    {
                        "tableId": "TABLE-001",
                        "rows": [
                            {
                                "rowIndex": 1,
                                "cells": [
                                    {"cellIndex": 1, "text": "固定表头"},
                                    {"cellIndex": 2, "text": "填写项"},
                                ],
                            }
                        ],
                    }
                ],
            }
            response_rows = [
                {
                    "sequence": 1,
                    "requirementId": "req-1",
                    "requirement": "技术要求",
                    "response": "按要求响应。",
                    "deviation": "待人工确认",
                    "remark": "",
                }
            ]

            output = assemble_docx(
                {
                    "title": "示例项目",
                    "plan": [
                        {
                            "sectionId": "tech-response",
                            "title": "技术和服务要求响应表",
                            "source": "copied_template",
                            "templateFragmentId": "tpl-unknown",
                            "templateResponseTable": {"mapped": False},
                            "requirementResponseRows": response_rows,
                            "content": "备用正文",
                        }
                    ],
                    "deviationTables": {"technical": [], "commercial": []},
                },
                root / "outputs",
                root=root,
                template_fragments=[fragment],
            )

            exported = Document(output["path"])
            all_cells = [
                cell.text
                for table in exported.tables
                for row in table.rows
                for cell in row.cells
            ]
            self.assertIn("固定表头", all_cells)
            self.assertIn("技术要求", all_cells)
            self.assertIn("按要求响应。", all_cells)
            self.assertTrue(output["exportWarnings"])


if __name__ == "__main__":
    unittest.main()
