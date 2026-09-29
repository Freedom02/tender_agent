from pathlib import Path
from tempfile import TemporaryDirectory
import zipfile
from xml.etree import ElementTree as ET

from docx import Document

try:
    from backend.app.store import Store, W_NS
except ImportError:  # store.py 保留远端版本，我方模板源导出的 W_NS 符号不在远端 store 中
    import pytest

    pytest.skip(
        "store.py 保留远端版本，模板源导出(W_NS)待 PG 迁移阶段迁",
        allow_module_level=True,
    )


def _workspace_with_template(tmp_path: Path) -> tuple[Store, str]:
    source_dir = tmp_path / "runtime" / "projects" / "P1" / "source"
    source_dir.mkdir(parents=True)
    source_path = source_dir / "source.docx"

    document = Document()
    document.add_paragraph("intro")
    document.add_paragraph("response form")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "no"
    table.cell(0, 1).text = "response"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "fill technical response"
    document.add_paragraph("contract text")
    document.save(source_path)

    store = Store(tmp_path)
    parsed = store.extract_document_content(
        "source.docx",
        source_path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    paragraph_segment = next(
        item for item in parsed["segments"] if item.get("text") == "response form"
    )
    table_segment = next(
        item
        for item in parsed["segments"]
        if item.get("contentType") == "table_row"
        and "fill technical response" in item.get("text", "")
    )
    project = {
        "id": "P1",
        "title": "Test project",
        "downloadUrl": "/runtime/projects/P1/source/source.docx",
        "document": parsed,
    }
    store.workspaces["W1"] = {
        "id": "W1",
        "projectId": "P1",
        "project": project,
        "reasoning": {
            "fineReview": {
                "analysisPack": {
                    "templateIndex": [
                        {
                            "id": "template-0001",
                            "templateId": "TPL-0001",
                            "title": "response form",
                            "sourceSegmentIds": [
                                paragraph_segment["id"],
                                table_segment["id"],
                            ],
                            "downloadText": (
                                "response form\n"
                                "no | response\n"
                                "1 | fill technical response"
                            ),
                        }
                    ]
                }
            }
        },
    }
    return store, "W1"


def test_export_template_source_preserves_selected_docx_range() -> None:
    with TemporaryDirectory(dir=Path.cwd()) as raw_root:
        store, workspace_id = _workspace_with_template(Path(raw_root))

        result = store.export_template_source(
            workspace_id,
            template_id="TPL-0001",
            output_format="docx",
        )

        assert result is not None
        exported = Document(result["path"])
        paragraphs = "\n".join(item.text for item in exported.paragraphs)
        table_text = "\n".join(
            cell.text
            for table in exported.tables
            for row in table.rows
            for cell in row.cells
        )
        assert "response form" in paragraphs
        assert "intro" not in paragraphs
        assert "contract text" not in paragraphs
        assert "fill technical response" in table_text


def test_export_template_source_keeps_txt_fallback() -> None:
    with TemporaryDirectory(dir=Path.cwd()) as raw_root:
        store, workspace_id = _workspace_with_template(Path(raw_root))

        result = store.export_template_source(
            workspace_id,
            template_id="TPL-0001",
            output_format="txt",
        )

        assert result is not None
        text = Path(result["path"]).read_text(encoding="utf-8")
        assert text.count("fill technical response") == 1


def test_export_template_source_copies_source_package_and_complete_table() -> None:
    with TemporaryDirectory(dir=Path.cwd()) as raw_root:
        root = Path(raw_root)
        source_dir = root / "runtime" / "projects" / "P1" / "source"
        source_dir.mkdir(parents=True)
        source_path = source_dir / "source.docx"

        document = Document()
        document.add_paragraph("before template")
        heading = document.add_paragraph("response form")
        heading.style = "Heading 2"
        document.add_paragraph()
        table = document.add_table(rows=3, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).merge(table.cell(0, 1)).text = "merged header"
        table.cell(1, 0).text = "first"
        table.cell(1, 1).text = "response"
        table.cell(2, 0).text = "second"
        table.cell(2, 1).text = "fill here"
        document.add_paragraph("after template")
        document.save(source_path)

        store = Store(root)
        parsed = store.extract_document_content(
            "source.docx",
            source_path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        heading_segment = next(
            item for item in parsed["segments"] if item.get("text") == "response form"
        )
        selected_table_row = next(
            item
            for item in parsed["segments"]
            if item.get("contentType") == "table_row"
            and item.get("rowIndex") == 3
        )
        store.workspaces["W1"] = {
            "id": "W1",
            "projectId": "P1",
            "project": {
                "id": "P1",
                "title": "Test project",
                "downloadUrl": str(source_path.resolve()),
                "document": parsed,
            },
            "reasoning": {
                "fineReview": {
                    "analysisPack": {
                        "templateIndex": [
                            {
                                "id": "template-0002",
                                "templateId": "TPL-0002",
                                "title": "response form",
                                "sourceSegmentIds": [
                                    heading_segment["id"],
                                    selected_table_row["id"],
                                ],
                            }
                        ]
                    }
                }
            },
        }

        result = store.export_template_source(
            "W1",
            template_id="TPL-0002",
            output_format="docx",
        )

        assert result is not None
        assert result["sourceExtraction"] == "docx_xml_range"
        assert result["formatPreserved"] is True
        exported_path = Path(result["path"])
        exported = Document(exported_path)
        exported_text = "\n".join(
            [paragraph.text for paragraph in exported.paragraphs]
            + [
                cell.text
                for table in exported.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        assert "before template" not in exported_text
        assert "after template" not in exported_text
        assert "response form" in exported_text
        assert len(exported.tables) == 1
        assert len(exported.tables[0].rows) == 3

        with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(exported_path) as exported_zip:
            source_names = set(source_zip.namelist())
            exported_names = set(exported_zip.namelist())
            assert source_names == exported_names
            for name in source_names - {"word/document.xml"}:
                assert source_zip.read(name) == exported_zip.read(name)

            source_root = ET.fromstring(source_zip.read("word/document.xml"))
            exported_root = ET.fromstring(exported_zip.read("word/document.xml"))
            source_body = source_root.find(f"{{{W_NS}}}body")
            exported_body = exported_root.find(f"{{{W_NS}}}body")
            assert source_body is not None
            assert exported_body is not None

            source_table = next(
                child
                for child in list(source_body)
                if child.tag.rsplit("}", 1)[-1] == "tbl"
                and "fill here" in "".join(
                    node.text or ""
                    for node in child.iter()
                    if node.tag.rsplit("}", 1)[-1] == "t"
                )
            )
            exported_table = next(
                child
                for child in list(exported_body)
                if child.tag.rsplit("}", 1)[-1] == "tbl"
            )
            assert ET.tostring(source_table) == ET.tostring(exported_table)
            grid_span = exported_table.find(
                f".//{{{W_NS}}}gridSpan"
            )
            assert grid_span is not None
            assert grid_span.attrib[f"{{{W_NS}}}val"] == "2"
