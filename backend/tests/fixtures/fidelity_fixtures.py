# -*- coding: utf-8 -*-
"""Deterministic local fixtures for the three bid-writing fidelity baselines.

The fixtures are created in a caller-owned temporary directory. No runtime
workspace data, network service, or binary template is required.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageDraw


def _grid_span(cell: Any) -> str:
    node = cell._tc.tcPr.find(qn("w:gridSpan"))
    return str(node.get(qn("w:val"))) if node is not None else "1"


def _v_merge(cell: Any) -> str:
    node = cell._tc.tcPr.find(qn("w:vMerge"))
    return str(node.get(qn("w:val"))) if node is not None else ""


def _paragraph_style(paragraph: Any) -> dict[str, Any]:
    run = paragraph.runs[0] if paragraph.runs else None
    size = run.font.size.pt if run and run.font.size else None
    return {
        "styleId": str(paragraph.style.style_id or ""),
        "alignment": str(paragraph.alignment) if paragraph.alignment is not None else "",
        "bold": bool(run.bold) if run else False,
        "fontSize": float(size) if size is not None else None,
    }


def _table_record(table: Any, body_index: int) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(table.rows, start=1):
        cells: list[dict[str, Any]] = []
        for cell_index, cell in enumerate(row.cells, start=1):
            cells.append(
                {
                    "cellIndex": cell_index,
                    "text": cell.text,
                    "gridSpan": _grid_span(cell),
                    "vMerge": _v_merge(cell),
                }
            )
        rows.append({"rowIndex": row_index, "cells": cells})
    return {
        "tableId": f"TABLE-{body_index:03d}",
        "bodyChildIndex": body_index,
        "rows": rows,
        "rowCount": len(rows),
        "columnCount": max((len(row["cells"]) for row in rows), default=0),
    }


def _document_blocks(document: Document) -> dict[str, Any]:
    """Build the fragment parser contract from the generated DOCX itself."""

    blocks: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    segment_ids: list[str] = []
    paragraph_number = 0
    table_number = 0
    for body_index, child in enumerate(document._element.body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            paragraph_number += 1
            paragraph = next(item for item in document.paragraphs if item._p is child)
            block_id = f"P-{paragraph_number:03d}"
            block = {
                "id": block_id,
                "text": paragraph.text,
                "contentType": "paragraph",
                "bodyChildIndex": body_index,
                "style": _paragraph_style(paragraph),
            }
            blocks.append(block)
            segment_ids.append(block_id)
        elif child.tag == qn("w:tbl"):
            table_number += 1
            table = next(item for item in document.tables if item._tbl is child)
            record = _table_record(table, body_index)
            record["tableId"] = f"TABLE-{table_number:03d}"
            tables.append(record)
            segment_ids.append(f"T-{table_number:03d}")

    root = blocks[0]
    return {
        "blocks": blocks,
        "units": blocks,
        "tables": tables,
        "sections": [
            {
                "id": "SEC-001",
                "sectionId": "SEC-001",
                "heading": root["text"],
                "level": 1,
                "startLocator": "fixture:format",
                "segmentIds": segment_ids,
            }
        ],
        "metadata": {"layout": {"pageWidthTwips": "12240", "marginLeftTwips": "1440"}},
    }


def create_template_fixture(root: Path) -> dict[str, Any]:
    """Create a styled, merged DOCX and its deterministic document-block view."""

    path = root / "fixture-template.docx"
    document = Document()
    root_heading = document.add_heading("第一章 投标文件格式", level=1)
    root_heading.runs[0].bold = True
    template_heading = document.add_paragraph("投标函格式")
    template_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    template_heading.style = document.styles["Heading 2"]
    template_heading.runs[0].bold = True
    template_heading.runs[0].font.size = Pt(16)
    document.add_paragraph("项目名称：________")
    document.add_paragraph("投标人名称：________")
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "固定表头"
    table.cell(0, 2).text = "填写项"
    table.cell(1, 0).text = "项目名称"
    table.cell(1, 1).text = "投标人名称"
    table.cell(1, 2).text = "________"
    table.cell(2, 0).text = "锁定文字"
    table.cell(2, 1).text = "内部编号"
    table.cell(2, 2).text = "________"
    document.add_paragraph("后续固定内容")
    document.add_heading("第二章 其他内容", level=1)
    document.add_paragraph("不属于投标函模板的章节")
    document.save(path)

    reread = Document(path)
    blocks = _document_blocks(reread)
    blocks["sections"][0]["segmentIds"] = [item["id"] for item in blocks["blocks"]] + [
        f"T-{index:03d}" for index in range(1, len(blocks["tables"]) + 1)
    ]
    return {"path": path, "documentBlocks": blocks}


def template_bundle(root: Path) -> dict[str, Any]:
    from backend.app.bid_writing.template_fragments import attach_template_fragments_to_bundle

    fixture = create_template_fixture(root)
    bundle = attach_template_fragments_to_bundle(
        {
            "projectId": "fixture-project",
            "workspaceId": "fixture-workspace",
            "sourceFile": {"filename": fixture["path"].name, "storagePath": fixture["path"].name},
            "documentBlocks": fixture["documentBlocks"],
            "bidFormatSpec": {
                "templates": [
                    {
                        "templateId": "tpl-fixture",
                        "title": "投标函格式",
                        "templateType": "fixed_template",
                        "sourceLocator": "投标函格式",
                        "editableFields": [
                            {"fieldId": "field-project", "label": "项目名称", "sourceKey": "projectName"},
                            {"fieldId": "field-company", "label": "投标人名称", "sourceKey": "companyName"},
                        ],
                    }
                ]
            },
        },
        root=root,
        source_path=fixture["path"],
    )
    return {**fixture, "bundle": bundle, "fragment": bundle["templateFragments"][0]}


def response_table_inputs() -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    section = {
        "sectionId": "response-table-01",
        "title": "技术和服务要求响应表",
        "responseType": "table",
        "itemType": "fixed_table",
    }
    fine_review = {
        "rows": [
            {
                "requirementId": "REQ-001",
                "category": "technical",
                "requirement": "平台支持统一身份认证。",
                "status": "满足",
                "reason": "产品资料已覆盖。",
                "matchedObjectIds": ["prod-001"],
            },
            {
                "requirementId": "REQ-002",
                "category": "service_plan",
                "requirement": "提供上线培训服务。",
                "status": "部分满足",
                "reason": "培训计划待补。",
                "gaps": ["培训课时"],
                "matchedObjectIds": [],
            },
            {"requirementId": "REQ-003", "category": "qualification", "requirement": "具有营业执照。", "status": "满足"},
        ]
    }
    knowledge = {
        "products": [
            {
                "id": "prod-001",
                "name": "统一认证模块",
                "description": "支持企业统一身份认证。",
                "evidencePath": "evidence/product-001.pdf",
                "verificationStatus": "verified",
            }
        ]
    }
    return section, fine_review, knowledge


def _write_scan_page(path: Path, label: str, fill: tuple[int, int, int]) -> Path:
    """Create a deterministic raster scan page without committing binary fixtures."""

    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (800, 500), fill)
    draw = ImageDraw.Draw(image)
    draw.rectangle((24, 24, 776, 476), outline=(25, 35, 45), width=4)
    draw.rectangle((48, 108, 752, 392), outline=(75, 90, 105), width=2)
    draw.text((64, 52), label, fill=(15, 25, 35))
    draw.text((72, 150), "PROGRAMMATIC SCAN PAGE", fill=(35, 45, 55))
    draw.text((72, 190), "FIDELITY-BASELINE / PAGE-01", fill=(35, 45, 55))
    image.save(path, format="PNG", optimize=False)
    return path


def evidence_fixture(root: Path) -> dict[str, Any]:
    """Create image-backed evidence and resolve it through the local media adapter."""

    evidence_dir = root / "evidence"
    evidence_dir.mkdir(parents=True, exist_ok=True)
    readable = evidence_dir / "01-product"
    second = evidence_dir / "02-certificate"
    unreadable = evidence_dir / "03-missing"
    product_page = _write_scan_page(
        readable / "page-01.png", "PRODUCT-PROOF-SCAN", (226, 239, 244)
    )
    certificate_page = _write_scan_page(
        second / "page-01.png", "CERTIFICATE-PROOF-SCAN", (242, 235, 218)
    )
    unreadable.mkdir(parents=True, exist_ok=True)
    (unreadable / ".keep").write_text("placeholder", encoding="utf-8")
    refs = {
        "readable": readable.as_posix(),
        "second": second.as_posix(),
        "unreadable": unreadable.as_posix(),
    }
    entries = [
        {
            "sequence": 1,
            "requirementId": "REQ-001",
            "requirement": "平台支持统一身份认证。",
            "status": "ready",
            "sourceRefs": [
                {"factId": "proof-01", "name": "产品证明", "evidencePath": refs["readable"]},
                {"factId": "proof-02", "name": "证书证明", "evidencePath": refs["second"]},
            ],
            "materials": [],
        },
        {
            "sequence": 2,
            "requirementId": "REQ-002",
            "requirement": "提供上线培训服务。",
            "status": "ready",
            "sourceRefs": [
                {"factId": "proof-01", "name": "产品证明", "evidencePath": refs["readable"]},
                {"factId": "proof-03", "name": "待补证明", "evidencePath": refs["unreadable"]},
            ],
            "materials": [],
        },
        {
            "sequence": 3,
            "requirementId": "REQ-003",
            "requirement": "提交培训计划。",
            "status": "missing",
            "sourceRefs": [],
            "materials": [],
        },
    ]
    from backend.app.bid_writing.evidence_assembler import assemble_evidence
    from backend.app.media_store.local import LocalFolderStore

    media_store = LocalFolderStore(root)
    source_refs = [
        ref
        for entry in entries
        for ref in entry["sourceRefs"]
    ]
    assets = assemble_evidence(source_refs, media_store)
    for asset in assets.values():
        asset["available"] = not bool(asset.get("placeholder"))
    # Keep the generated page paths available to byte-level DOCX assertions.
    assets[refs["readable"]]["fixture_page"] = str(product_page)
    assets[refs["second"]]["fixture_page"] = str(certificate_page)
    return {"entries": entries, "assets": assets, "refs": refs}
